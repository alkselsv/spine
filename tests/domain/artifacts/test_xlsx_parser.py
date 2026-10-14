from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
from uuid import UUID
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from spine.domain.artifacts import ParseCoverage, XlsxCellLocator
from spine.domain.sources import RevisionMetadata, SourceRevision, content_revision
from spine.domain.sources.original_reference import ObjectReference


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
WORKSPACE = UUID(int=1)
ENVIRONMENT = UUID(int=2)
SOURCE_OBJECT = UUID(int=3)
REVISION = UUID(int=4)


def revision(content: bytes, *, workspace_id: UUID = WORKSPACE) -> SourceRevision:
    digest = sha256(content).hexdigest()
    return content_revision(
        revision_id=REVISION,
        source_object_id=SOURCE_OBJECT,
        workspace_id=workspace_id,
        environment_id=ENVIRONMENT,
        kind="content",
        revision_schema_version="source-revision:v1",
        revision_metadata_schema="revision-metadata:r1-document-v1",
        revision_metadata=RevisionMetadata(embedded_title="fixture"),
        original_reference=ObjectReference(
            schema_version=1,
            object_id=UUID(int=5),
            storage_generation=UUID(int=6),
            digest_algorithm="sha256",
            digest_hex=digest,
            byte_length=len(content),
        ),
        original_sha256=digest,
        byte_length=len(content),
        media_type=XLSX_MEDIA_TYPE,
        observed_at=datetime(2026, 10, 13, tzinfo=timezone.utc),
    )


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Prices"
    sheet["A1"] = "Product"
    sheet["A3"] = "repeat"
    sheet["A5"] = "repeat"
    sheet.merge_cells("A9:B9")
    sheet["A9"] = "merged"
    second = workbook.create_sheet("Other")
    second["C4"] = "other"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def workbook_with_cached_and_uncached_formulas() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "=1+1"
    sheet["A2"] = "=2+2"
    sheet["A3"] = "='[external.xlsx]Sheet1'!A1"
    output = BytesIO()
    workbook.save(output)
    original = output.getvalue()
    rewritten = BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(rewritten, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            value = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = value.decode("utf-8")
                xml = xml.replace('<c r="A1"><f>1+1</f><v></v></c>', '<c r="A1"><f>1+1</f><v>2</v></c>')
                value = xml.encode("utf-8")
            target.writestr(item, value)
    return rewritten.getvalue()


def test_xlsx_parser_preserves_physical_coordinates_and_resolves_exact_cells() -> None:
    from spine.ingestion.parsers import XlsxParser, resolve_xlsx_locator

    content = workbook_bytes()
    document = XlsxParser().parse(revision(content), content, media_type=XLSX_MEDIA_TYPE)

    assert document.coverage is ParseCoverage.COMPLETE
    assert [(segment.text, segment.locator) for segment in document.segments] == [
        ("Product", XlsxCellLocator(sheet="Prices", row=1, cell_range="A1")),
        ("repeat", XlsxCellLocator(sheet="Prices", row=3, cell_range="A3")),
        ("repeat", XlsxCellLocator(sheet="Prices", row=5, cell_range="A5")),
        ("merged", XlsxCellLocator(sheet="Prices", row=9, cell_range="A9:B9")),
        ("other", XlsxCellLocator(sheet="Other", row=4, cell_range="C4")),
    ]
    assert document.segments[1].locator != document.segments[2].locator
    assert resolve_xlsx_locator(document, content, document.segments[3].locator) == "merged"


def test_parse_document_selects_xlsx_adapter_without_using_legacy_loader() -> None:
    from spine.ingestion.parsers import parse_document

    content = workbook_bytes()
    document = parse_document(revision(content), content)

    assert document.parser_name == "xlsx"
    assert document.locator_contract_version == "xlsx-cell:v1"


def test_xlsx_parser_excludes_uncached_formulas_without_disclosing_formula_text() -> None:
    from spine.ingestion.parsers import XlsxParser

    content = workbook_with_cached_and_uncached_formulas()
    document = XlsxParser().parse(revision(content), content, media_type=XLSX_MEDIA_TYPE)

    assert document.coverage is ParseCoverage.PARTIAL
    assert [segment.text for segment in document.segments] == ["2"]
    assert {item.code for item in document.exclusions} >= {"formula_without_cached_value", "external_link"}
    assert all(
        "1+1" not in item.model_dump_json() and "2+2" not in item.model_dump_json()
        for item in document.exclusions
    )


def test_xlsx_parser_excludes_hidden_and_unsupported_content() -> None:
    from spine.ingestion.parsers import XlsxParser

    workbook = Workbook()
    visible = workbook.active
    visible["A1"] = "visible"
    visible["A2"] = "hidden row"
    visible.row_dimensions[2].hidden = True
    visible["B1"] = "hidden column"
    visible.column_dimensions["B"].hidden = True
    visible["A3"] = 1
    visible["B3"] = 2
    chart = BarChart()
    chart.add_data(Reference(visible, min_col=1, min_row=3, max_col=2, max_row=3))
    visible.add_chart(chart, "D1")
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "secret"
    output = BytesIO()
    workbook.save(output)
    content = output.getvalue()

    document = XlsxParser().parse(revision(content), content, media_type=XLSX_MEDIA_TYPE)

    assert document.coverage is ParseCoverage.PARTIAL
    assert [segment.text for segment in document.segments] == ["visible", "1"]
    assert {item.code for item in document.exclusions} >= {"hidden_content", "unsupported_structure"}
    assert "secret" not in document.model_dump_json()
    assert "hidden row" not in document.model_dump_json()


@pytest.mark.parametrize(
    ("content", "parser_kwargs", "code"),
    [
        (b"not an xlsx", {}, "malformed_structure"),
        (workbook_bytes(), {"max_bytes": 8}, "bounded_input"),
        (workbook_bytes(), {"max_cells": 2}, "bounded_input"),
    ],
)
def test_xlsx_parser_reports_malformed_and_bounded_inputs(
    content: bytes, parser_kwargs: dict[str, int], code: str
) -> None:
    from spine.ingestion.parsers import XlsxParser

    document = XlsxParser(**parser_kwargs).parse(revision(content), content, media_type=XLSX_MEDIA_TYPE)

    assert code in {item.code for item in document.exclusions}
    assert document.coverage in {ParseCoverage.FAILED, ParseCoverage.PARTIAL}


def test_xlsx_locator_round_trip_rejects_changed_original_bytes() -> None:
    from spine.ingestion.parsers import XlsxParser, resolve_xlsx_locator

    content = workbook_bytes()
    document = XlsxParser().parse(revision(content), content, media_type=XLSX_MEDIA_TYPE)

    with pytest.raises(ValueError, match="original content does not match source revision"):
        resolve_xlsx_locator(document, content + b"changed", document.segments[0].locator)


@pytest.mark.asyncio
async def test_xlsx_artifact_repository_rejects_revision_parser_and_workspace_mismatches() -> None:
    from spine.application.persistence.parsed_artifacts import ParsedArtifactCommand
    from spine.infrastructure.persistence.parsed_artifacts import InMemoryParsedArtifactRepository
    from spine.ingestion.parsers import XlsxParser

    content = workbook_bytes()
    document = XlsxParser().parse(revision(content), content, media_type=XLSX_MEDIA_TYPE)
    repository = InMemoryParsedArtifactRepository()
    await repository.persist(ParsedArtifactCommand.create(document, command_identity="parse-xlsx"))
    locator = document.segments[0].locator

    assert await repository.resolve_segment(
        workspace_id=document.workspace_id,
        environment_id=document.environment_id,
        source_revision_id=document.source_revision_id,
        source_revision_digest=document.source_revision_digest,
        original_sha256=document.original_sha256,
        parser_version=document.parser_version,
        locator_contract_version=document.locator_contract_version,
        locator=locator,
    ) == document.segments[0]
    assert await repository.resolve_segment(
        workspace_id=document.workspace_id,
        environment_id=document.environment_id,
        source_revision_id=document.source_revision_id,
        source_revision_digest="f" * 64,
        original_sha256=document.original_sha256,
        parser_version=document.parser_version,
        locator_contract_version=document.locator_contract_version,
        locator=locator,
    ) is None
    assert await repository.resolve_segment(
        workspace_id=UUID(int=7),
        environment_id=document.environment_id,
        source_revision_id=document.source_revision_id,
        source_revision_digest=document.source_revision_digest,
        original_sha256=document.original_sha256,
        parser_version=document.parser_version,
        locator_contract_version=document.locator_contract_version,
        locator=locator,
    ) is None
    assert await repository.resolve_segment(
        workspace_id=document.workspace_id,
        environment_id=document.environment_id,
        source_revision_id=UUID(int=8),
        source_revision_digest=document.source_revision_digest,
        original_sha256=document.original_sha256,
        parser_version=document.parser_version,
        locator_contract_version=document.locator_contract_version,
        locator=locator,
    ) is None
    assert await repository.resolve_segment(
        workspace_id=document.workspace_id,
        environment_id=document.environment_id,
        source_revision_id=document.source_revision_id,
        source_revision_digest=document.source_revision_digest,
        original_sha256=document.original_sha256,
        parser_version="other-parser:v1",
        locator_contract_version=document.locator_contract_version,
        locator=locator,
    ) is None
    assert await repository.resolve_segment(
        workspace_id=document.workspace_id,
        environment_id=document.environment_id,
        source_revision_id=document.source_revision_id,
        source_revision_digest=document.source_revision_digest,
        original_sha256=document.original_sha256,
        parser_version=document.parser_version,
        locator_contract_version="other-locator:v1",
        locator=locator,
    ) is None
