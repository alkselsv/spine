from __future__ import annotations

from io import BytesIO
from uuid import UUID

from docx import Document
from docx.oxml import OxmlElement
import pytest

from spine.application.persistence.errors import ParsedArtifactConflictError
from spine.application.persistence.parsed_artifacts import ParsedArtifactCommand
from spine.domain.artifacts import (
    ExtractedOoxmlText,
    OoxmlBlockPathStep,
    OoxmlTextLocator,
    ParseCoverage,
)
from spine.infrastructure.persistence.parsed_artifacts import InMemoryParsedArtifactRepository
from spine.ingestion.parsers import DocxParser, parse_document
from tests.domain.artifacts.test_parsed_documents import revision


DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def docx_bytes(document: Document) -> bytes:
    target = BytesIO()
    document.save(target)
    return target.getvalue()


def test_docx_nested_tables_and_duplicate_text_have_exact_structural_paths() -> None:
    source_document = Document()
    source_document.add_paragraph("repeat")
    table = source_document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "repeat"
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "nested"
    content = docx_bytes(source_document)

    document = parse_document(revision(content, media_type=DOCX_MEDIA_TYPE), content)

    assert document.coverage is ParseCoverage.COMPLETE
    assert [segment.text for segment in document.segments] == ["repeat", "repeat", "nested"]
    assert [segment.locator for segment in document.segments] == [
        OoxmlTextLocator(
            block_path=(OoxmlBlockPathStep(kind="paragraph", index=0),),
            start_codepoint=0,
            end_codepoint=6,
        ),
        OoxmlTextLocator(
            block_path=(
                OoxmlBlockPathStep(kind="table", index=1),
                OoxmlBlockPathStep(kind="row", index=0),
                OoxmlBlockPathStep(kind="cell", index=0),
                OoxmlBlockPathStep(kind="paragraph", index=1),
            ),
            start_codepoint=0,
            end_codepoint=6,
        ),
        OoxmlTextLocator(
            block_path=(
                OoxmlBlockPathStep(kind="table", index=1),
                OoxmlBlockPathStep(kind="row", index=0),
                OoxmlBlockPathStep(kind="cell", index=1),
                OoxmlBlockPathStep(kind="table", index=2),
                OoxmlBlockPathStep(kind="row", index=0),
                OoxmlBlockPathStep(kind="cell", index=0),
                OoxmlBlockPathStep(kind="paragraph", index=1),
            ),
            start_codepoint=0,
            end_codepoint=6,
        ),
    ]
    assert document.segments[0].segment_id != document.segments[1].segment_id
    assert document.parser_name == "docx"
    assert document.locator_contract_version == "docx-ooxml-block-codepoint:v1"


def test_docx_unsupported_story_is_a_content_free_partial_exclusion() -> None:
    source_document = Document()
    source_document.add_paragraph("visible body")
    source_document.sections[0].header.paragraphs[0].text = "protected header text"
    content = docx_bytes(source_document)

    document = parse_document(revision(content, media_type=DOCX_MEDIA_TYPE), content)

    assert document.coverage is ParseCoverage.PARTIAL
    assert [segment.text for segment in document.segments] == ["visible body"]
    assert [(item.code, item.count, item.locator) for item in document.exclusions] == [
        ("unsupported_story", 1, None),
    ]
    assert "protected header text" not in document.model_dump_json()


def test_docx_drawing_is_a_typed_partial_exclusion() -> None:
    source_document = Document()
    paragraph = source_document.add_paragraph("visible text")
    drawing = OxmlElement("w:drawing")
    protected = OxmlElement("w:t")
    protected.text = "protected drawing text"
    drawing.append(protected)
    paragraph.add_run()._r.append(drawing)
    content = docx_bytes(source_document)

    document = parse_document(revision(content, media_type=DOCX_MEDIA_TYPE), content)

    assert document.coverage is ParseCoverage.PARTIAL
    assert [segment.text for segment in document.segments] == ["visible text"]
    assert [(item.code, item.count, item.locator) for item in document.exclusions] == [
        (
            "unsupported_drawing",
            1,
            OoxmlTextLocator(
                block_path=(OoxmlBlockPathStep(kind="paragraph", index=0),),
                start_codepoint=0,
                end_codepoint=12,
            ),
        ),
    ]
    assert "protected drawing text" not in document.model_dump_json()


@pytest.mark.parametrize("drawing_tag", ["w:pict", "w:object"])
def test_docx_legacy_drawing_is_a_typed_partial_exclusion(drawing_tag: str) -> None:
    source_document = Document()
    paragraph = source_document.add_paragraph("visible text")
    drawing = OxmlElement(drawing_tag)
    protected = OxmlElement("w:t")
    protected.text = "protected legacy drawing text"
    drawing.append(protected)
    paragraph.add_run()._r.append(drawing)
    content = docx_bytes(source_document)

    document = parse_document(revision(content, media_type=DOCX_MEDIA_TYPE), content)

    assert document.coverage is ParseCoverage.PARTIAL
    assert [segment.text for segment in document.segments] == ["visible text"]
    assert [item.code for item in document.exclusions] == ["unsupported_drawing"]
    assert "protected legacy drawing text" not in document.model_dump_json()


def test_docx_merged_cell_ambiguity_has_failed_coverage() -> None:
    source_document = Document()
    table = source_document.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "protected merged text"
    content = docx_bytes(source_document)

    document = parse_document(revision(content, media_type=DOCX_MEDIA_TYPE), content)

    assert document.coverage is ParseCoverage.FAILED
    assert document.segments == ()
    assert [(item.code, item.count, item.locator) for item in document.exclusions] == [
        ("ambiguous_merged_cell", 1, None),
    ]
    assert "protected merged text" not in document.model_dump_json()


def test_docx_unknown_body_block_has_content_free_failed_coverage() -> None:
    source_document = Document()
    unknown = OxmlElement("w:sdt")
    protected = OxmlElement("w:t")
    protected.text = "protected unknown text"
    unknown.append(protected)
    source_document.element.body.insert(0, unknown)
    content = docx_bytes(source_document)

    document = parse_document(revision(content, media_type=DOCX_MEDIA_TYPE), content)

    assert document.coverage is ParseCoverage.FAILED
    assert document.segments == ()
    assert [(item.code, item.count, item.locator) for item in document.exclusions] == [
        ("unsupported_structure", 1, None),
    ]
    assert "protected unknown text" not in document.model_dump_json()


def test_docx_unknown_table_block_has_content_free_failed_coverage() -> None:
    source_document = Document()
    table = source_document.add_table(rows=1, cols=1)
    unknown = OxmlElement("w:sdt")
    protected = OxmlElement("w:t")
    protected.text = "protected table structure"
    unknown.append(protected)
    table._tbl.append(unknown)
    content = docx_bytes(source_document)

    document = parse_document(revision(content, media_type=DOCX_MEDIA_TYPE), content)

    assert document.coverage is ParseCoverage.FAILED
    assert document.segments == ()
    assert [item.code for item in document.exclusions] == ["unsupported_structure"]
    assert "protected table structure" not in document.model_dump_json()


def test_malformed_docx_has_content_free_failed_coverage() -> None:
    content = b"not a docx and protected customer text"

    document = parse_document(revision(content, media_type=DOCX_MEDIA_TYPE), content)

    assert document.coverage is ParseCoverage.FAILED
    assert document.segments == ()
    assert [(item.code, item.count, item.locator) for item in document.exclusions] == [
        ("malformed_document", 1, None),
    ]
    assert b"protected customer text" not in document.model_dump_json().encode()


def test_docx_byte_limit_has_bounded_content_free_exclusion() -> None:
    source_document = Document()
    source_document.add_paragraph("protected customer text")
    content = docx_bytes(source_document)

    document = DocxParser(max_bytes=1).parse(
        revision(content, media_type=DOCX_MEDIA_TYPE),
        content,
    )

    assert document.coverage is ParseCoverage.FAILED
    assert document.segments == ()
    assert [(item.code, item.count, item.locator) for item in document.exclusions] == [
        ("bounded_input", 1, None),
    ]
    assert "protected customer text" not in document.model_dump_json()
    assert DocxParser(max_bytes=1).config_reference != DocxParser(max_bytes=2).config_reference


@pytest.mark.parametrize("limit", ["entries", "uncompressed_bytes"])
def test_docx_archive_limit_has_bounded_content_free_exclusion(limit: str) -> None:
    source_document = Document()
    source_document.add_paragraph("protected customer text")
    content = docx_bytes(source_document)
    parser = (
        DocxParser(max_archive_entries=1)
        if limit == "entries"
        else DocxParser(max_uncompressed_bytes=1)
    )

    document = parser.parse(revision(content, media_type=DOCX_MEDIA_TYPE), content)

    assert document.coverage is ParseCoverage.FAILED
    assert document.segments == ()
    assert [item.code for item in document.exclusions] == ["bounded_input"]
    assert "protected customer text" not in document.model_dump_json()


def test_docx_nesting_limit_excludes_deeper_tables_without_leaking_text() -> None:
    source_document = Document()
    source_document.add_paragraph("visible body")
    table = source_document.add_table(rows=1, cols=1)
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "protected nested text"
    content = docx_bytes(source_document)

    document = DocxParser(max_nesting_depth=1).parse(
        revision(content, media_type=DOCX_MEDIA_TYPE),
        content,
    )

    assert document.coverage is ParseCoverage.PARTIAL
    assert [segment.text for segment in document.segments] == ["visible body"]
    assert [item.code for item in document.exclusions] == ["bounded_input"]
    assert "protected nested text" not in document.model_dump_json()


@pytest.mark.parametrize(
    ("parser", "expected_text"),
    [
        (DocxParser(max_blocks=1), ["safe"]),
        (DocxParser(max_paragraph_chars=3), []),
        (DocxParser(max_total_chars=5), ["safe"]),
    ],
    ids=["blocks", "paragraph-characters", "total-characters"],
)
def test_docx_structure_limit_produces_bounded_coverage(
    parser: DocxParser,
    expected_text: list[str],
) -> None:
    source_document = Document()
    source_document.add_paragraph("safe")
    source_document.add_paragraph("protected customer text")
    content = docx_bytes(source_document)

    document = parser.parse(revision(content, media_type=DOCX_MEDIA_TYPE), content)

    assert [segment.text for segment in document.segments] == expected_text
    assert document.coverage is (
        ParseCoverage.PARTIAL if expected_text else ParseCoverage.FAILED
    )
    assert {item.code for item in document.exclusions} == {"bounded_input"}
    assert "protected customer text" not in document.model_dump_json()


def test_docx_block_limit_stops_traversal_with_one_bounded_exclusion() -> None:
    source_document = Document()
    source_document.add_paragraph("safe")
    for index in range(100):
        source_document.add_paragraph(f"protected-{index}")
    content = docx_bytes(source_document)

    document = DocxParser(max_blocks=1).parse(
        revision(content, media_type=DOCX_MEDIA_TYPE),
        content,
    )

    assert [segment.text for segment in document.segments] == ["safe"]
    assert [(item.code, item.count) for item in document.exclusions] == [
        ("bounded_input", 1),
    ]
    assert "protected-99" not in document.model_dump_json()


def test_docx_block_limit_counts_unknown_structure_and_aggregates_output() -> None:
    source_document = Document()
    for index in range(100):
        unknown = OxmlElement("w:sdt")
        protected = OxmlElement("w:t")
        protected.text = f"protected-{index}"
        unknown.append(protected)
        source_document.element.body.insert(index, unknown)
    content = docx_bytes(source_document)

    document = DocxParser(max_blocks=1).parse(
        revision(content, media_type=DOCX_MEDIA_TYPE),
        content,
    )

    assert document.coverage is ParseCoverage.FAILED
    assert document.segments == ()
    assert [(item.code, item.count) for item in document.exclusions] == [
        ("unsupported_structure", 1),
        ("bounded_input", 1),
    ]
    assert "protected-99" not in document.model_dump_json()


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("source_revision_id", UUID(int=99)),
        ("source_revision_digest", "f" * 64),
        ("original_sha256", "f" * 64),
        ("parser_version", "docx:v2"),
        ("locator_contract_version", "docx-ooxml-block-codepoint:v2"),
    ],
)
def test_docx_locator_round_trip_rejects_artifact_identity_mismatch(
    field: str,
    value: object,
) -> None:
    source_document = Document()
    source_document.add_paragraph("repeat")
    source_document.add_paragraph("repeat")
    content = docx_bytes(source_document)
    source = revision(content, media_type=DOCX_MEDIA_TYPE)
    parser = DocxParser()
    document = parser.parse(source, content)
    locator = document.segments[1].locator

    assert document.resolve_locator(
        ExtractedOoxmlText(block_path=locator.block_path, text="repeat"),
        locator,
    ) == "repeat"
    assert parser.resolve_locator(source, content, document, locator) == "repeat"
    with pytest.raises(ValueError, match="locator does not resolve"):
        parser.resolve_locator(
            source,
            content,
            document.model_copy(update={field: value}),
            locator,
        )


@pytest.mark.asyncio
async def test_docx_artifact_replay_has_one_immutable_effect() -> None:
    source_document = Document()
    source_document.add_paragraph("replay text")
    content = docx_bytes(source_document)
    document = parse_document(revision(content, media_type=DOCX_MEDIA_TYPE), content)
    command = ParsedArtifactCommand.create(document, command_identity="docx-parse-1")
    repository = InMemoryParsedArtifactRepository()

    first = await repository.persist(command)
    replay = await repository.persist(command)
    resolved = await repository.resolve_segment(
        workspace_id=document.workspace_id,
        environment_id=document.environment_id,
        source_revision_id=document.source_revision_id,
        source_revision_digest=document.source_revision_digest,
        original_sha256=document.original_sha256,
        parser_version=document.parser_version,
        locator_contract_version=document.locator_contract_version,
        locator=document.segments[0].locator,
    )

    assert first.replay is False
    assert replay.replay is True
    assert resolved == document.segments[0]
    assert await repository.count() == 1


@pytest.mark.asyncio
async def test_docx_artifact_conflict_rolls_back_without_mutation() -> None:
    first_source = Document()
    first_source.add_paragraph("first")
    first_content = docx_bytes(first_source)
    second_source = Document()
    second_source.add_paragraph("second")
    second_content = docx_bytes(second_source)
    first = parse_document(revision(first_content, media_type=DOCX_MEDIA_TYPE), first_content)
    second = parse_document(revision(second_content, media_type=DOCX_MEDIA_TYPE), second_content)
    repository = InMemoryParsedArtifactRepository()

    await repository.persist(
        ParsedArtifactCommand.create(first, command_identity="docx-parse-1")
    )
    with pytest.raises(ParsedArtifactConflictError):
        await repository.persist(
            ParsedArtifactCommand.create(second, command_identity="docx-parse-1")
        )

    assert await repository.count() == 1
    assert await repository.get(
        workspace_id=first.workspace_id,
        environment_id=first.environment_id,
        artifact_id=first.artifact_id,
    ) == first


@pytest.mark.asyncio
async def test_docx_artifacts_are_isolated_between_two_workspaces() -> None:
    source_document = Document()
    source_document.add_paragraph("isolated text")
    content = docx_bytes(source_document)
    first = parse_document(revision(content, media_type=DOCX_MEDIA_TYPE), content)
    second = parse_document(
        revision(
            content,
            media_type=DOCX_MEDIA_TYPE,
            workspace_id=UUID(int=7),
            environment_id=UUID(int=8),
        ),
        content,
    )
    repository = InMemoryParsedArtifactRepository()

    await repository.persist(
        ParsedArtifactCommand.create(first, command_identity="docx-parse-1")
    )
    await repository.persist(
        ParsedArtifactCommand.create(second, command_identity="docx-parse-1")
    )

    assert await repository.count() == 2
    assert await repository.get(
        workspace_id=second.workspace_id,
        environment_id=second.environment_id,
        artifact_id=first.artifact_id,
    ) is None
