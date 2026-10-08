"""PROTOTYPE ONLY: reproducible locator experiment for GitHub Issue #13.

Run with ``uv run python tests/prototypes/issue_13_stable_locators.py``.
All documents are synthetic and generated in a temporary directory.
"""

from __future__ import annotations

import hashlib
import json
import re
import tempfile
from datetime import datetime
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from spine.ingest.loaders import load_path as legacy_load_path


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def freeze_zip(path: Path) -> None:
    """Remove creation-time ZIP headers from generated synthetic fixtures."""
    target = path.with_suffix(path.suffix + ".fixed")
    with ZipFile(path) as source, ZipFile(target, "w") as output:
        for filename in source.namelist():
            info = ZipInfo(filename, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            content = source.read(filename)
            if path.suffix == ".xlsx" and filename == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                                 rb"\g<1>2000-01-01T00:00:00Z\g<2>", content)
            output.writestr(info, content)
    target.replace(path)


def record(revision: str, original_sha256: str, kind: str, locator: dict, raw: str) -> dict:
    # The revision is supplied by the caller; a filename or digest cannot identify
    # a SourceObject or decide whether a revision is canonically current.
    return {
        "source_revision_id": revision,
        "original_sha256": original_sha256,
        "kind": kind,
        "locator": locator,
        "raw": raw,
        "normalized": " ".join(raw.split()),
    }


def text_records(path: Path, revision: str, markdown: bool) -> list[dict]:
    original_sha256 = digest(path)
    source = path.read_bytes().decode("utf-8", errors="strict")
    out = []
    offset = 0
    headings: list[str] = []
    for number, line in enumerate(source.splitlines(keepends=True), 1):
        content = line.rstrip("\r\n")
        if markdown and content.startswith("#"):
            level = len(content) - len(content.lstrip("#"))
            if level <= 6 and content[level:level + 1] == " ":
                headings = headings[: level - 1] + [content[level + 1:]]
        if content.strip():
            out.append(record(revision, original_sha256, "markdown" if markdown else "text", {
                "line": number,
                "char_start": offset,
                "char_end": offset + len(content),
                "heading_path": headings.copy() if markdown else [],
            }, content))
        offset += len(line)
    return out


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_paragraph_text(paragraph) -> str:
    pieces = []
    for node in paragraph.iter():
        if node.tag == W + "drawing":
            raise ValueError("DOCX drawing has no citation-grade text locator")
        if node.tag == W + "t":
            pieces.append(node.text or "")
        elif node.tag == W + "tab":
            pieces.append("\t")
        elif node.tag in (W + "br", W + "cr"):
            pieces.append("\n")
    return "".join(pieces)


def docx_records(path: Path, revision: str) -> list[dict]:
    original_sha256 = digest(path)
    document = Document(path)
    unsupported_stories = {"header", "footer", "footnotes", "endnotes", "comments"}
    for relationship in document.part.rels.values():
        story = relationship.reltype.rsplit("/", 1)[-1]
        if story in unsupported_stories:
            raise ValueError(f"DOCX {story} story needs a locator strategy")
    out = []

    def walk(parent, prefix: list[dict]) -> None:
        for index, child in enumerate(parent):
            if child.tag == W + "sectPr" and not prefix:
                continue
            if child.tag == W + "p":
                raw = docx_paragraph_text(child)
                if raw.strip():
                    out.append(record(revision, original_sha256, "docx", {
                        "block_path": prefix + [{"block": index, "type": "paragraph"}],
                        "char_start": 0,
                        "char_end": len(raw),
                    }, raw))
            elif child.tag == W + "tbl":
                for row_index, row in enumerate(child.findall(W + "tr")):
                    for cell_index, cell in enumerate(row.findall(W + "tc")):
                        if cell.find(".//" + W + "gridSpan") is not None or cell.find(".//" + W + "vMerge") is not None:
                            raise ValueError("DOCX merged table cell needs grid-aware locator")
                        walk(cell, prefix + [{"block": index, "type": "table"},
                                             {"row": row_index, "cell": cell_index}])
            elif child.tag not in (W + "tcPr",):
                raise ValueError(f"DOCX unsupported block {child.tag.rsplit('}', 1)[-1]}")

    walk(document.element.body, [])
    if not out:
        raise ValueError("DOCX has no extractable text")
    return out


def xlsx_records(path: Path, revision: str) -> list[dict]:
    original_sha256 = digest(path)
    workbook = load_workbook(path, read_only=False, data_only=True)
    formulas = load_workbook(path, read_only=False, data_only=False)
    out = []
    for sheet in workbook.worksheets:
        formula_sheet = formulas[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                formula_cell = formula_sheet[cell.coordinate]
                if formula_cell.data_type == "f" and cell.value is None:
                    raise ValueError(f"XLSX formula {sheet.title}!{cell.coordinate} has no cached value")
                if cell.value is None:
                    continue
                raw = str(cell.value)
                if raw.strip():
                    merged = next((str(span) for span in sheet.merged_cells.ranges
                                   if cell.coordinate in span), None)
                    out.append(record(revision, original_sha256, "xlsx", {
                        "sheet": sheet.title,
                        "row": cell.row,
                        "cell": cell.coordinate,
                        "cell_range": merged or cell.coordinate,
                    }, raw))
    if not out:
        raise ValueError("XLSX has no extractable values")
    return out


def pdf_records(path: Path, revision: str) -> list[dict]:
    original_sha256 = digest(path)
    out = []
    for page_number, page in enumerate(PdfReader(path).pages, 1):
        extracted = page.extract_text() or ""
        if not extracted.strip():
            raise ValueError(f"PDF page {page_number} has no extractable text")
        # These offsets address pypdf's extracted page text, not PDF bytes or
        # visual coordinates. They are useful only with a pinned parser version.
        offset = 0
        for line in extracted.splitlines(keepends=True):
            content = line.rstrip("\r\n")
            start = offset
            offset += len(line)
            if not content.strip():
                continue
            out.append(record(revision, original_sha256, "pdf", {
                "page": page_number,
                "extracted_char_start": start,
                "extracted_char_end": start + len(content),
            }, content))
    return out


def make_pdf(path: Path, *, blank: bool = False) -> None:
    writer = PdfWriter()
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    font_ref = writer._add_object(font)
    for page_index, lines in enumerate(() if blank else
                                       (("Repeat value", "Repeat value"), ("Item", "Value", "A", "42"))):
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})})
        stream = DecodedStreamObject()
        commands = []
        for index, line in enumerate(lines):
            if page_index == 1:
                x = 72 if index % 2 == 0 else 300
                y = 720 if index < 2 else 690
            else:
                x, y = 72, 720 - 14 * index
            commands.append(f"BT /F1 12 Tf {x} {y} Td ({line}) Tj ET")
        if page_index == 1:
            commands.append("72 680 m 400 680 l 236 680 m 236 730 l S")
        stream.set_data(" ".join(commands).encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    if blank:
        writer.add_blank_page(width=612, height=792)
    with path.open("wb") as target:
        writer.write(target)


def fixtures(directory: Path) -> dict[str, Path]:
    paths = {extension: directory / f"sample.{extension}"
             for extension in ("md", "txt", "docx", "xlsx", "pdf")}
    paths["md"].write_bytes("# Один\r\nПовтор 😀\r\n\r\n## Таблица\r\n| A | B |\r\n| Повтор | 42 |\r\n".encode())
    paths["txt"].write_bytes("Повтор 😀\n\nПовтор 😀\n".encode())
    doc = Document()
    doc.add_paragraph("Повтор 😀")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Повтор 😀"
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Вложено"
    doc.add_paragraph("Конец")
    doc.save(paths["docx"])
    freeze_zip(paths["docx"])
    book = Workbook()
    book.properties.created = datetime(2000, 1, 1)
    book.properties.modified = datetime(2000, 1, 1)
    sheet = book.active
    sheet.title = "Данные"
    sheet["A1"] = "Название"
    sheet["A3"] = "Повтор 😀"
    sheet["B5"] = "Повтор 😀"
    sheet.merge_cells("A9:B9")
    sheet["A9"] = "Общее"
    book.save(paths["xlsx"])
    freeze_zip(paths["xlsx"])
    make_pdf(paths["pdf"])
    return paths


PARSERS = {"md": lambda p, r: text_records(p, r, True),
           "txt": lambda p, r: text_records(p, r, False),
           "docx": docx_records, "xlsx": xlsx_records, "pdf": pdf_records}


def resolve_excerpt(path: Path, item: dict, expected_revision: str) -> str:
    """Reopen the original and resolve the proposed locator independently."""
    if item["source_revision_id"] != expected_revision:
        raise ValueError("SourceRevision identity mismatch")
    if item["original_sha256"] != digest(path):
        raise ValueError("SourceRevision original digest mismatch")
    kind = item["kind"]
    where = item["locator"]
    if kind in ("markdown", "text"):
        return path.read_bytes().decode("utf-8")[where["char_start"]:where["char_end"]]
    if kind == "pdf":
        extracted = PdfReader(path).pages[where["page"] - 1].extract_text() or ""
        return extracted[where["extracted_char_start"]:where["extracted_char_end"]]
    if kind == "xlsx":
        value = load_workbook(path, data_only=True)[where["sheet"]][where["cell"]].value
        return "" if value is None else str(value)
    if kind == "docx":
        node = Document(path).element.body
        for part in where["block_path"]:
            if "block" in part:
                node = node[part["block"]]
            else:
                node = node.findall(W + "tr")[part["row"]].findall(W + "tc")[part["cell"]]
        text = docx_paragraph_text(node)
        return text[where["char_start"]:where["char_end"]]
    raise AssertionError(kind)


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="spine-issue-13-") as name:
        root = Path(name)
        paths = fixtures(root)
        findings = {}
        for extension, path in paths.items():
            revision = f"synthetic-revision-{extension}-v1"
            first = PARSERS[extension](path, revision)
            assert first == PARSERS[extension](path, revision)
            assert all(item["source_revision_id"] == revision for item in first)
            assert len({json.dumps(item["locator"], sort_keys=True) for item in first}) == len(first)
            assert all(resolve_excerpt(path, item, revision) == item["raw"] for item in first)
            assert all(" ".join(item["raw"].split()) == item["normalized"] for item in first)
            findings[extension] = {"sha256": digest(path), "records": first}
        assert [r["locator"]["row"] for r in findings["xlsx"]["records"]] == [1, 3, 5, 9]
        legacy_xlsx_paths = [item.source_path.rsplit("!", 1)[-1]
                             for item in legacy_load_path(paths["xlsx"])]
        assert legacy_xlsx_paths == ["row2", "row4", "row8"]
        findings["legacy_xlsx_rows"] = legacy_xlsx_paths
        assert findings["xlsx"]["records"][-1]["locator"]["cell_range"] == "A9:B9"
        assert {"row": 0, "cell": 1} in findings["docx"]["records"][2]["locator"]["block_path"]
        assert findings["md"]["records"][1]["raw"] == "Повтор 😀"
        assert findings["txt"]["records"][0]["locator"] != findings["txt"]["records"][1]["locator"]
        assert findings["pdf"]["records"][0]["locator"] != findings["pdf"]["records"][1]["locator"]
        for extension in ("md", "txt"):
            source = paths[extension].read_bytes().decode("utf-8")
            for item in findings[extension]["records"]:
                where = item["locator"]
                assert source[where["char_start"]:where["char_end"]] == item["raw"]
        old_txt = findings["txt"]["records"]
        old_txt_path = root / "old.txt"
        old_txt_path.write_bytes(paths["txt"].read_bytes())
        paths["txt"].write_bytes("Новое\n".encode())
        new_txt = text_records(paths["txt"], "synthetic-revision-txt-v2", False)
        assert new_txt[0]["source_revision_id"] != old_txt[0]["source_revision_id"]
        assert new_txt[0]["raw"] != old_txt[0]["raw"]
        assert resolve_excerpt(old_txt_path, old_txt[0], "synthetic-revision-txt-v1") == old_txt[0]["raw"]
        try:
            resolve_excerpt(old_txt_path, old_txt[0], "synthetic-revision-txt-v2")
        except ValueError as exc:
            findings["revision_identity_mismatch"] = str(exc)
        else:
            raise AssertionError("Wrong revision identity was accepted")
        same_excerpt_path = root / "same-excerpt.txt"
        same_excerpt_path.write_bytes("Повтор 😀\n\nПовтор 😀\nдругой хвост\n".encode())
        try:
            resolve_excerpt(same_excerpt_path, old_txt[0], "synthetic-revision-txt-v1")
        except ValueError as exc:
            findings["revision_digest_mismatch"] = str(exc)
        else:
            raise AssertionError("Different original with same excerpt was accepted")
        findings["revision_change"] = {"old": old_txt[0], "new": new_txt[0]}
        formula_path = root / "formula.xlsx"
        book = Workbook()
        book.properties.created = datetime(2000, 1, 1)
        book.properties.modified = datetime(2000, 1, 1)
        book.active["C7"] = "=1+1"
        book.save(formula_path)
        freeze_zip(formula_path)
        try:
            xlsx_records(formula_path, "formula-revision")
        except ValueError as exc:
            findings["xlsx_uncached_formula"] = str(exc)
        else:
            raise AssertionError("Uncached formula was accepted")
        merged_docx_path = root / "merged.docx"
        merged_doc = Document()
        merged_table = merged_doc.add_table(rows=1, cols=2)
        merged_table.cell(0, 0).merge(merged_table.cell(0, 1)).text = "Общее"
        merged_doc.save(merged_docx_path)
        freeze_zip(merged_docx_path)
        try:
            docx_records(merged_docx_path, "merged-revision")
        except ValueError as exc:
            findings["docx_merged_cell"] = str(exc)
        else:
            raise AssertionError("Merged DOCX cell was accepted")
        header_docx_path = root / "header.docx"
        header_doc = Document()
        header_doc.add_paragraph("Body")
        header_doc.sections[0].header.paragraphs[0].text = "Header evidence"
        header_doc.save(header_docx_path)
        freeze_zip(header_docx_path)
        try:
            docx_records(header_docx_path, "header-revision")
        except ValueError as exc:
            findings["docx_header"] = str(exc)
        else:
            raise AssertionError("DOCX header text was silently omitted")
        blank_pdf = root / "blank.pdf"
        make_pdf(blank_pdf, blank=True)
        try:
            pdf_records(blank_pdf, "blank-revision")
        except ValueError as exc:
            findings["pdf_no_text"] = str(exc)
        else:
            raise AssertionError("Blank PDF was accepted")
        invalid_text = root / "invalid.txt"
        invalid_text.write_bytes(b"\xff")
        try:
            text_records(invalid_text, "invalid-revision", False)
        except UnicodeDecodeError:
            findings["invalid_utf8"] = "rejected"
        else:
            raise AssertionError("Invalid UTF-8 was accepted")
        for extension in ("pdf", "docx", "xlsx"):
            invalid = root / f"invalid.{extension}"
            invalid.write_bytes(b"not a document")
            try:
                PARSERS[extension](invalid, "invalid-revision")
            except Exception as exc:
                findings[f"invalid_{extension}"] = type(exc).__name__
            else:
                raise AssertionError(f"Invalid {extension} was accepted")
        print(json.dumps(findings, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
