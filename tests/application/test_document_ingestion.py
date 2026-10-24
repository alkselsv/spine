from __future__ import annotations

import hashlib
from collections.abc import AsyncIterable, AsyncIterator
from datetime import datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock
from uuid import NAMESPACE_URL, UUID, uuid5

from docx import Document
from openpyxl import Workbook
import pytest
from pydantic import ValidationError

from spine.auth import AuthorizedRequestContext, RequestedAuthorizationScope
from spine.application.document_ingestion import (
    DocumentIngestionOrchestrator,
    DocumentIngestionStatus,
    DocumentUploadCommand,
    ProjectionRequestPayload,
    build_projection_request_intent,
)
from spine.evals.corpus_preparation import (
    ApprovedCorpusArtifact,
    ApprovedCorpusPreparationService,
)
from spine.evals.real_qa_r1 import (
    REAL_R1_DATASET_DIGEST,
    REAL_R1_DATASET_VERSION,
    RealRuntimeMappingManifest,
    RealRuntimeRevisionIntent,
    RealRuntimeSourceIntent,
)
from spine.application.ingestion import BoundedContentValidator
from spine.application.persistence.context import PersistenceOperation
from spine.application.persistence.idempotency import OpaqueResultReference
from spine.application.persistence.outbox import OutboxEventRegistry
from spine.application.persistence.parsed_artifacts import ParsedArtifactPersistenceResult
from spine.application.persistence.ingestion import IngestionMutationResult
from spine.domain.ingestion import (
    IngestionItemState,
    IngestionRunState,
    IngestionTargetKind,
    IngestionTransition,
    ProcessingReceipt,
)
from spine.domain.sources import AdmissionDisposition
from spine.infrastructure.object_storage import (
    IntegrityResult,
    IntegrityStatus,
    ObjectReference,
    ObservedContentIdentity,
    OriginalUploadCommandResult,
    PreWriteRequestDigest,
    WriteReceipt,
)
from spine.infrastructure.persistence.parsed_artifacts import InMemoryParsedArtifactRepository
from spine.ingestion.parsers import (
    DOCX_MEDIA_TYPE,
    XLSX_MEDIA_TYPE,
    DocxParser,
    PdfParser,
    XlsxParser,
    parse_document,
)
from spine.application.ingestion_faults import (
    IngestionFaultInjected,
    IngestionFaultReason,
)
from spine.infrastructure.ingestion_fault_injection import (
    IngestionFaultBoundary,
    IngestionFaultController,
    IngestionFaultInjectionProfile,
    IngestionFaultRule,
)
from tests.ingestion.test_pdf_parser import pdf_bytes


NOW = datetime(2026, 10, 14, 12, 0, tzinfo=timezone.utc)
WORKSPACE = UUID("10000000-0000-4000-8000-000000000001")
ENVIRONMENT = UUID("20000000-0000-4000-8000-000000000001")


def _id(value: str) -> UUID:
    return uuid5(NAMESPACE_URL, f"document-ingestion-test:{value}")


async def _stream(*parts: bytes) -> AsyncIterator[bytes]:
    for part in parts:
        yield part


class _Uploader:
    def __init__(self, content: bytes) -> None:
        digest = hashlib.sha256(content).hexdigest()
        reference = ObjectReference(
            schema_version=1, object_id=_id("object"), storage_generation=_id("generation"),
            digest_algorithm="sha256", digest_hex=digest, byte_length=len(content),
        )
        receipt = WriteReceipt(
            schema_version=1, upload_id=_id("upload"), object_reference=reference,
            observed_content=ObservedContentIdentity(observed_sha256=digest, observed_byte_length=len(content)),
            verification=IntegrityResult(
                status=IntegrityStatus.VERIFIED, algorithm="sha256",
                observed_digest=digest, observed_length=len(content),
            ), storage_status="finalized_verified", finalized_at=NOW,
        )
        result = OriginalUploadCommandResult(
            request_digest=PreWriteRequestDigest(
                algorithm="sha256", algorithm_version="spine.command-digest.v1", digest_hex="a" * 64,
            ), upload_id=receipt.upload_id, object_id=reference.object_id,
            issue8_result_reference={"receipt_id": _id("issue8-receipt")},
        )
        self.outcome = SimpleNamespace(result=result, write_receipt=receipt, replay=False)
        self.calls = 0

    async def upload(self, context: object, command: object, chunks: AsyncIterable[bytes]):
        del context, command, chunks
        self.calls += 1
        return self.outcome


class _Reader:
    def __init__(self, content: bytes) -> None:
        self.content = content
        self.calls = 0

    def read_finalized(self, **kwargs: object) -> AsyncIterable[bytes]:
        assert kwargs["workspace_id"] == WORKSPACE
        assert kwargs["environment_id"] == ENVIRONMENT
        self.calls += 1
        return _stream(self.content[:2], self.content[2:])


class _Ingestion:
    def __init__(self) -> None:
        self.run = None
        self.item = None
        self.transitions: list[IngestionTransition] = []
        self.binding = None
        self.receipt: ProcessingReceipt | None = None
        self.failures = []

    def _head(self, *, kind: IngestionTargetKind, run_id: UUID, item_id: UUID | None = None):
        candidates = [
            value for value in self.transitions
            if value.target_kind is kind and value.ingestion_run_id == run_id
            and value.ingestion_run_source_id == item_id
        ]
        return candidates[-1] if candidates else None

    async def create_run(self, command):
        replay = self.run is not None
        if self.run is None:
            self.run = command.run
            self.transitions.append(IngestionTransition(
                ingestion_transition_id=_id("run-root"), workspace_id=command.run.workspace_id,
                environment_id=command.run.environment_id, target_kind=IngestionTargetKind.RUN,
                ingestion_run_id=command.run.ingestion_run_id, state=IngestionRunState.REQUESTED,
                attempt=1, actor_or_service="fake", trace_id=_id("trace"), occurred_at=NOW,
            ))
        return SimpleNamespace(value=self.run, replay=replay)

    async def create_item(self, item):
        if self.item is None:
            self.item = item
            self.transitions.append(IngestionTransition(
                ingestion_transition_id=_id("item-root"), workspace_id=item.workspace_id,
                environment_id=item.environment_id, target_kind=IngestionTargetKind.ITEM,
                ingestion_run_id=item.ingestion_run_id, ingestion_run_source_id=item.ingestion_run_source_id,
                state=IngestionItemState.RECEIVED, attempt=1, actor_or_service="fake",
                trace_id=_id("trace"), occurred_at=NOW,
            ))
        elif self.item != item:
            raise AssertionError("Ingestion item identity conflicts.")
        return self.item

    async def resolve_run_head(self, run_id):
        return self._head(kind=IngestionTargetKind.RUN, run_id=run_id)

    async def resolve_item_head(self, item_id):
        return self._head(kind=IngestionTargetKind.ITEM, run_id=self.item.ingestion_run_id, item_id=item_id)

    async def bind_source(self, command):
        if self.binding is None:
            self.binding = command.binding
            head = await self.resolve_item_head(command.item.ingestion_run_source_id)
            self.transitions.append(IngestionTransition(
                ingestion_transition_id=_id("observed"), workspace_id=command.item.workspace_id,
                environment_id=command.item.environment_id, target_kind=IngestionTargetKind.ITEM,
                ingestion_run_id=command.item.ingestion_run_id,
                ingestion_run_source_id=command.item.ingestion_run_source_id,
                predecessor_transition_id=head.ingestion_transition_id, state=IngestionItemState.OBSERVED,
                attempt=head.attempt + 1, actor_or_service="fake", trace_id=_id("trace"), occurred_at=NOW,
            ))
        return SimpleNamespace(value=self.binding, replay=self.binding == command.binding)

    async def append_transition(self, transition):
        self.transitions.append(transition)
        return transition

    async def resolve_processing_receipt(self, item_id):
        return self.receipt

    async def resolve_failures(self, item_id):
        return tuple(
            failure
            for failure in self.failures
            if failure.ingestion_run_source_id == item_id
        )

    async def append_warning(self, warning):
        return warning

    async def append_failure(self, failure):
        self.failures.append(failure)
        return failure

    async def record_processing_receipt(self, command):
        replay = self.receipt is not None
        self.receipt = self.receipt or command.receipt
        return SimpleNamespace(value=self.receipt, replay=replay)

    async def resolve_item(self, item_id):
        return self.item

    def snapshot(self):
        return (
            self.run,
            self.item,
            list(self.transitions),
            self.binding,
            self.receipt,
        )

    def restore(self, snapshot) -> None:
        self.run = snapshot[0]
        self.item = snapshot[1]
        self.transitions = list(snapshot[2])
        self.binding = snapshot[3]
        self.receipt = snapshot[4]


class _Admissions:
    def __init__(self) -> None:
        self.decision = None

    async def resolve_effective(self, revision_id):
        return self.decision

    async def append_decision(self, decision):
        self.decision = decision
        return decision


class _Canonical:
    def __init__(self) -> None:
        self.calls = 0
        self.current_revision_id: UUID | None = None

    async def accept(self, command):
        self.calls += 1
        self.current_revision_id = command.revision.revision_id
        return command

    async def resolve_current(self, source_object_id):
        del source_object_id
        if self.current_revision_id is None:
            return None
        return SimpleNamespace(current_revision_id=self.current_revision_id)


class _Uow:
    def __init__(self, ingestion, admissions, canonical, outbox):
        self.ingestion = ingestion
        self.admissions = admissions
        self.canonical = canonical
        self.outbox = outbox
        self.trace_id = _id("trace")
        self._ingestion_snapshot = ingestion.snapshot()

    async def __aenter__(self):
        return self

    async def __aexit__(self, *args):
        if args[0] is not None:
            await self.rollback()
        return None

    async def commit(self):
        try:
            self.outbox.commit()
        except Exception:
            self.ingestion.restore(self._ingestion_snapshot)
            raise
        self._ingestion_snapshot = self.ingestion.snapshot()

    async def rollback(self):
        self.ingestion.restore(self._ingestion_snapshot)
        self.outbox.rollback()

    def _trace_id_snapshot(self):
        return self.trace_id


class _Outbox:
    def __init__(self) -> None:
        self.events = []
        self.pending = []
        self.fail_commit_with_event_type: str | None = None

    async def append(self, intent):
        self.pending.append(intent)
        return intent.event_id

    def commit(self) -> None:
        if any(
            event.event_type == self.fail_commit_with_event_type
            for event in self.pending
        ):
            self.pending.clear()
            raise RuntimeError("injected transaction failure")
        self.events.extend(self.pending)
        self.pending.clear()

    def rollback(self) -> None:
        self.pending.clear()


def _orchestrator(
    content: bytes,
    *,
    media_type: str,
    parser_calls: list[int],
    pre_acceptance_faults: list[bool] | None = None,
    fault_injector=None,
    clock=lambda: NOW,
):
    ingestion = _Ingestion()
    admissions = _Admissions()
    canonical = _Canonical()
    outbox = _Outbox()
    uow = _Uow(ingestion, admissions, canonical, outbox)
    uploader = _Uploader(content)
    reader = _Reader(content)

    def parse(revision, value: bytes, *, media_type: str):
        parser_calls.append(1)
        return parse_document(revision, value, media_type=media_type)

    def context_factory(operation: PersistenceOperation, trace_id: UUID):
        return SimpleNamespace(operation=operation, trace_id=trace_id)

    def canonical_command_factory(**kwargs):
        assert kwargs["request"].access_policy_version_id == _id("policy")
        if pre_acceptance_faults and pre_acceptance_faults.pop():
            raise IngestionFaultInjected(
                IngestionFaultReason.PRE_ACCEPTANCE_RETRYABLE
            )
        return SimpleNamespace(revision=kwargs["revision"])

    class PolicyEnsurer:
        async def ensure(self, *, request, source, context, activated_at):
            del request, source, context, activated_at
            return _id("policy")

    return DocumentIngestionOrchestrator(
        uow_factory=lambda context: uow,
        uploader=uploader,
        validation_reader=reader,
        validator=BoundedContentValidator(max_bytes=10_000_000, max_chunk_bytes=10_000_000),
        parsed_artifacts=InMemoryParsedArtifactRepository(),
        context_factory=context_factory,
        access_policy_ensurer=PolicyEnsurer(),
        canonical_command_factory=canonical_command_factory,
        outbox_events=OutboxEventRegistry(),
        parser=parse,
        fault_injector=fault_injector,
        clock=clock,
        max_validation_bytes=10_000_000,
        max_parse_bytes=10_000_000,
    ), uploader, reader, canonical, outbox, ingestion


@pytest.mark.asyncio
async def test_pre_acceptance_fault_is_visible_and_retry_does_not_create_early_current_revision() -> None:
    parser_calls: list[int] = []
    orchestrator, _, _, canonical, _, ingestion = _orchestrator(
        b"qualification source\n",
        media_type="text/plain",
        parser_calls=parser_calls,
        pre_acceptance_faults=[True],
    )

    failed = await orchestrator.ingest(
        request=_request("text/plain"), chunks=_stream(b"qualification source\n")
    )

    assert failed.status is DocumentIngestionStatus.FAILED
    assert failed.failure_ids
    assert canonical.current_revision_id is None
    assert ingestion.failures[0].code == "pre_acceptance_retryable"
    assert ingestion.failures[0].retryable is True

    retried = await orchestrator.ingest(
        request=_request("text/plain"), chunks=_stream(b"qualification source\n")
    )

    assert retried.status is DocumentIngestionStatus.SUCCEEDED
    assert retried.revision.revision_id == failed.revision.revision_id
    assert canonical.current_revision_id == retried.revision.revision_id


@pytest.mark.parametrize(
    ("boundary", "reason"),
    (
        (IngestionFaultBoundary.PARSER, IngestionFaultReason.PARSER_RETRYABLE),
        (
            IngestionFaultBoundary.ORIGINAL_STORAGE,
            IngestionFaultReason.ORIGINAL_STORAGE_UNAVAILABLE,
        ),
    ),
)
async def test_processing_fault_is_reported_with_safe_reason_code(
    boundary: IngestionFaultBoundary,
    reason: IngestionFaultReason,
) -> None:
    command = _request("text/plain")
    controller = IngestionFaultController(
        IngestionFaultInjectionProfile(
            profile_id="qualification-ingestion-test",
            profile_version="r1",
            deployment_profile_id="qualification-isolated-test",
            deployment_profile_version="r1",
            deployment_fingerprint="a" * 64,
            rules=(
                IngestionFaultRule(
                    workspace_id=WORKSPACE,
                    environment_id=ENVIRONMENT,
                    command_identity=command.command_identity,
                    boundary=boundary,
                    reason_code=reason,
                ),
            ),
        )
    )
    orchestrator, _, reader, canonical, _, ingestion = _orchestrator(
        b"qualification source\n",
        media_type="text/plain",
        parser_calls=[],
        fault_injector=controller,
    )

    result = await orchestrator.ingest(
        request=command, chunks=_stream(b"qualification source\n")
    )

    assert result.status is DocumentIngestionStatus.FAILED
    assert result.processing_receipt is not None
    assert result.failure_ids
    assert ingestion.failures[0].code == reason.value
    assert ingestion.failures[0].retryable is True
    assert canonical.current_revision_id == result.revision.revision_id
    if boundary is IngestionFaultBoundary.ORIGINAL_STORAGE:
        assert reader.calls == 1


def _request(media_type: str, *, key: str = "document-1") -> DocumentUploadCommand:
    filename = {
        "text/plain": "document.txt",
        "text/markdown": "document.md",
        "application/pdf": "document.pdf",
        DOCX_MEDIA_TYPE: "document.docx",
        XLSX_MEDIA_TYPE: "document.xlsx",
    }.get(media_type, "document.bin")
    return DocumentUploadCommand(
        workspace_id=WORKSPACE, environment_id=ENVIRONMENT,
        command_identity=(
            "api.document.upload."
            + hashlib.sha256(key.encode("utf-8")).hexdigest()[:32]
        ), idempotency_key=key,
        upload_identity="uploaded-document", purpose="document.upload.v1",
        declared_media_type=media_type, declared_filename=filename,
    )


def _r1_document(media_type: str) -> bytes:
    if media_type == "text/plain":
        return b"hello\nworld\n"
    if media_type == "text/markdown":
        return b"# heading\n\nmarkdown body\n"
    if media_type == "application/pdf":
        return pdf_bytes(("visible PDF text",))
    if media_type == DOCX_MEDIA_TYPE:
        document = Document()
        document.add_paragraph("visible DOCX text")
        target = BytesIO()
        document.save(target)
        return target.getvalue()
    if media_type == XLSX_MEDIA_TYPE:
        workbook = Workbook()
        workbook.active["A1"] = "visible XLSX text"
        target = BytesIO()
        workbook.save(target)
        return target.getvalue()
    raise AssertionError(f"unhandled media type: {media_type}")


class _ApprovedCorpusArtifacts:
    def __init__(self, artifact: ApprovedCorpusArtifact) -> None:
        self._artifact = artifact

    async def resolve(
        self, *, dataset_source_id: str, dataset_revision_id: str
    ) -> ApprovedCorpusArtifact | None:
        if (
            dataset_source_id == self._artifact.dataset_source_id
            and dataset_revision_id == self._artifact.dataset_revision_id
        ):
            return self._artifact
        return None


@pytest.mark.asyncio
async def test_approved_corpus_preparation_uses_canonical_ingestion_and_outbox() -> None:
    content = b"hello\nworld\n"
    digest = hashlib.sha256(content).hexdigest()
    manifest = RealRuntimeMappingManifest(
        dataset_version=REAL_R1_DATASET_VERSION,
        dataset_sha256=REAL_R1_DATASET_DIGEST,
        owner_approval_ref="approval:integration-test",
        sources=tuple(
            RealRuntimeSourceIntent(
                dataset_source_id=f"source:approved-{index}",
                upload_identity=f"approved.source.{index}",
                revisions=(
                    RealRuntimeRevisionIntent(
                        dataset_revision_id=f"revision:approved-{index}",
                        expected_original_sha256=(
                            digest
                            if index == 0
                            else hashlib.sha256(f"missing-{index}".encode()).hexdigest()
                        ),
                        ingestion_intent_id=f"intent:approved-{index}",
                    ),
                ),
            )
            for index in range(30)
        ),
    )
    artifact = ApprovedCorpusArtifact(
        dataset_source_id="source:approved-0",
        dataset_revision_id="revision:approved-0",
        original_sha256=digest,
        byte_length=len(content),
        media_type="text/plain",
        filename="approved-0.txt",
        content=content,
    )
    parser_calls: list[int] = []
    orchestrator, _, _, canonical, outbox, _ = _orchestrator(
        content, media_type="text/plain", parser_calls=parser_calls
    )
    context = MagicMock(spec=AuthorizedRequestContext)
    context.scope = RequestedAuthorizationScope(
        workspace_id=WORKSPACE, environment_id=ENVIRONMENT
    )
    context.acting_subject_id = _id("operator")
    context.authorization_generation = 1

    report = await ApprovedCorpusPreparationService(
        artifacts=_ApprovedCorpusArtifacts(artifact), ingestion=orchestrator
    ).prepare(manifest=manifest, trusted_context=context)

    assert report.revisions[0].status == "prepared"
    assert all(item.status == "missing" for item in report.revisions[1:])
    assert canonical.calls == 1
    assert parser_calls == [1]
    assert [event.event_type for event in outbox.events] == [
        "ingestion.processing.requested",
        "ingestion.projection.requested",
    ]


@pytest.mark.asyncio
async def test_document_ingestion_completes_text_path_and_replays_stably() -> None:
    parser_calls: list[int] = []
    orchestrator, uploader, reader, canonical, outbox, ingestion = _orchestrator(
        b"hello\nworld\n", media_type="text/plain", parser_calls=parser_calls
    )
    first = await orchestrator.ingest(request=_request("text/plain"), chunks=_stream(b"hello\nworld\n"))
    second = await orchestrator.ingest(request=_request("text/plain"), chunks=_stream(b"hello\nworld\n"))

    assert first.status is DocumentIngestionStatus.SUCCEEDED
    assert first.artifact is not None
    assert first.processing_receipt is not None
    assert second.replay is True
    assert second.processing_receipt == first.processing_receipt
    assert len(parser_calls) == 1
    assert uploader.calls == 2
    assert reader.calls == 3
    assert canonical.calls == 1
    assert [event.event_type for event in outbox.events] == [
        "ingestion.processing.requested",
        "ingestion.projection.requested",
    ]
    projection_request = outbox.events[1]
    assert projection_request.payload_json() == {
        "schema_version": "ingestion-projection-request:v1",
        "ingestion_run_id": str(first.run.ingestion_run_id),
        "ingestion_run_source_id": str(first.item.ingestion_run_source_id),
        "source_revision_id": str(first.revision.revision_id),
        "processing_receipt_id": str(first.processing_receipt.processing_receipt_id),
        "processor_kind": "document-ingestion",
        "parser_version": "text-markdown:v1",
        "parser_config_reference": "text-markdown-default:v1",
        "locator_contract_version": "physical-line-codepoint:v1",
    }
    assert projection_request.correlation_id == first.run.ingestion_run_id
    assert projection_request.causation_id == outbox.events[0].event_id
    replayed_request = build_projection_request_intent(
        registry=orchestrator.outbox_events,
        run=first.run,
        item=first.item,
        receipt=first.processing_receipt,
        trace_id=projection_request.trace_id,
        causation_id=outbox.events[0].event_id,
    )
    changed_digest_request = build_projection_request_intent(
        registry=orchestrator.outbox_events,
        run=first.run,
        item=first.item,
        receipt=first.processing_receipt.model_copy(
            update={"original_sha256": "f" * 64}
        ),
        trace_id=projection_request.trace_id,
        causation_id=outbox.events[0].event_id,
    )
    assert replayed_request.event_id == projection_request.event_id
    assert changed_digest_request.event_id != projection_request.event_id
    assert (
        changed_digest_request.producer_deduplication_id
        == projection_request.producer_deduplication_id
    )
    assert (
        await ingestion.resolve_item_head(first.item.ingestion_run_source_id)
    ).state is IngestionItemState.PROJECTION_REQUESTED
    item_states = [
        transition.state
        for transition in ingestion.transitions
        if transition.ingestion_run_source_id == first.item.ingestion_run_source_id
    ]
    assert item_states[-2:] == [
        IngestionItemState.PROCESSED,
        IngestionItemState.PROJECTION_REQUESTED,
    ]


@pytest.mark.asyncio
async def test_document_ingestion_replay_reuses_item_when_clock_advances() -> None:
    clock_now = [NOW]
    orchestrator, _, _, _, _, _ = _orchestrator(
        b"hello\nworld\n",
        media_type="text/plain",
        parser_calls=[],
        clock=lambda: clock_now[0],
    )
    request = _request("text/plain")

    first = await orchestrator.ingest(
        request=request, chunks=_stream(b"hello\nworld\n")
    )
    clock_now[0] = NOW.replace(minute=NOW.minute + 1)
    replay = await orchestrator.ingest(
        request=request, chunks=_stream(b"hello\nworld\n")
    )

    assert replay.replay is True
    assert replay.item == first.item


@pytest.mark.asyncio
async def test_replacement_ingestion_preserves_existing_source_creation_time() -> None:
    source_created_at = NOW - timedelta(days=1)
    orchestrator, *_ = _orchestrator(
        b"replacement\n",
        media_type="text/plain",
        parser_calls=[],
    )
    request = _request("text/plain", key="replacement").model_copy(
        update={"source_created_at": source_created_at}
    )

    result = await orchestrator.ingest(
        request=request,
        chunks=_stream(b"replacement\n"),
    )

    assert result.source.created_at == source_created_at
    assert result.item.created_at == NOW


@pytest.mark.parametrize(
    "media_type",
    ("text/plain", "text/markdown", "application/pdf", DOCX_MEDIA_TYPE, XLSX_MEDIA_TYPE),
)
@pytest.mark.asyncio
async def test_document_ingestion_completes_each_r1_format_with_parser_metadata(
    media_type: str,
) -> None:
    content = _r1_document(media_type)
    parser_calls: list[int] = []
    orchestrator, _, _, _, outbox, _ = _orchestrator(
        content, media_type=media_type, parser_calls=parser_calls
    )

    result = await orchestrator.ingest(
        request=_request(media_type), chunks=_stream(content)
    )

    assert result.status is DocumentIngestionStatus.SUCCEEDED
    assert result.artifact is not None
    assert result.processing_receipt is not None
    assert result.processing_receipt.parser_name == result.artifact.parser_name
    assert result.processing_receipt.parser_version == result.artifact.parser_version
    assert result.processing_receipt.locator_contract_version == (
        result.artifact.locator_contract_version
    )
    assert result.processing_receipt.output_reference == result.artifact.artifact_id
    segment = result.artifact.segments[0]
    if media_type in {"text/plain", "text/markdown"}:
        assert result.artifact.resolve_locator(content.decode(), segment.locator) == segment.text
    elif media_type == "application/pdf":
        assert PdfParser().resolve_locator(result.revision, content, result.artifact, segment.locator) == segment.text
    elif media_type == DOCX_MEDIA_TYPE:
        assert DocxParser().resolve_locator(result.revision, content, result.artifact, segment.locator) == segment.text
    else:
        assert XlsxParser().resolve_locator(result.artifact, content, segment.locator) == segment.text
    assert len(parser_calls) == 1
    assert [event.event_type for event in outbox.events] == [
        "ingestion.processing.requested",
        "ingestion.projection.requested",
    ]


@pytest.mark.asyncio
async def test_document_ingestion_marks_unknown_format_unsupported_without_parser_access() -> None:
    parser_calls: list[int] = []
    orchestrator, _, _, _, outbox, _ = _orchestrator(
        b"opaque", media_type="application/octet-stream", parser_calls=parser_calls
    )
    result = await orchestrator.ingest(
        request=_request("application/octet-stream"), chunks=_stream(b"opaque")
    )

    assert result.status is DocumentIngestionStatus.UNSUPPORTED
    assert result.processing_receipt is not None
    assert result.processing_receipt.coverage.value == "failed"
    assert parser_calls == []
    assert outbox.events == []


@pytest.mark.asyncio
async def test_document_ingestion_failed_parse_keeps_parser_evidence_without_output_artifact() -> None:
    content = pdf_bytes(())
    orchestrator, _, _, _, outbox, _ = _orchestrator(
        content, media_type="application/pdf", parser_calls=[]
    )

    result = await orchestrator.ingest(
        request=_request("application/pdf"), chunks=_stream(content)
    )

    assert result.status is DocumentIngestionStatus.FAILED
    assert result.artifact is not None
    assert result.processing_receipt is not None
    assert result.processing_receipt.parser_name == "pdf"
    assert result.processing_receipt.output_reference is None
    assert result.processing_receipt.output_digest is None
    assert result.processing_receipt.exclusions[0].code == "invalid_locator"
    assert result.processing_receipt.failure_ids
    assert outbox.events == []

async def test_projection_request_is_not_visible_when_its_transaction_fails() -> None:
    parser_calls: list[int] = []
    orchestrator, _, _, _, outbox, ingestion = _orchestrator(
        b"hello\n", media_type="text/plain", parser_calls=parser_calls
    )
    outbox.fail_commit_with_event_type = "ingestion.projection.requested"

    with pytest.raises(RuntimeError, match="injected transaction failure"):
        await orchestrator.ingest(
            request=_request("text/plain"),
            chunks=_stream(b"hello\n"),
        )

    assert outbox.events == []
    assert ingestion.receipt is None
    assert (
        await ingestion.resolve_item_head(ingestion.item.ingestion_run_source_id)
    ).state is IngestionItemState.PROCESSING

    outbox.fail_commit_with_event_type = None
    recovered = await orchestrator.ingest(
        request=_request("text/plain"),
        chunks=_stream(b"hello\n"),
    )

    assert recovered.replay is True
    assert recovered.processing_receipt is not None
    assert [event.event_type for event in outbox.events] == [
        "ingestion.processing.requested",
        "ingestion.projection.requested",
    ]
    assert (
        await ingestion.resolve_item_head(recovered.item.ingestion_run_source_id)
    ).state is IngestionItemState.PROJECTION_REQUESTED


@pytest.mark.asyncio
async def test_stale_revision_is_processed_without_projection_request() -> None:
    parser_calls: list[int] = []
    orchestrator, _, _, canonical, outbox, ingestion = _orchestrator(
        b"hello\n", media_type="text/plain", parser_calls=parser_calls
    )

    async def resolve_newer_current(source_object_id: UUID):
        del source_object_id
        return SimpleNamespace(current_revision_id=_id("newer-revision"))

    canonical.resolve_current = resolve_newer_current
    result = await orchestrator.ingest(
        request=_request("text/plain"),
        chunks=_stream(b"hello\n"),
    )

    assert result.status is DocumentIngestionStatus.SUCCEEDED
    assert [event.event_type for event in outbox.events] == [
        "ingestion.processing.requested"
    ]
    assert (
        await ingestion.resolve_item_head(result.item.ingestion_run_source_id)
    ).state is IngestionItemState.PROCESSED


@pytest.mark.asyncio
@pytest.mark.parametrize("mismatch", ("workspace", "lineage"))
async def test_projection_request_rejects_mismatched_scope_or_lineage(
    mismatch: str,
) -> None:
    parser_calls: list[int] = []
    orchestrator, _, _, _, outbox, _ = _orchestrator(
        b"hello\n", media_type="text/plain", parser_calls=parser_calls
    )
    result = await orchestrator.ingest(
        request=_request("text/plain"),
        chunks=_stream(b"hello\n"),
    )
    assert result.processing_receipt is not None
    item = result.item
    receipt = result.processing_receipt
    if mismatch == "workspace":
        item = item.model_copy(update={"workspace_id": _id("other-workspace")})
    else:
        receipt = receipt.model_copy(
            update={"ingestion_run_source_id": _id("other-item")}
        )
    event_count = len(outbox.events)

    with pytest.raises(
        ValueError,
        match="projection request ingestion scope or lineage is invalid",
    ) as captured:
        build_projection_request_intent(
            registry=orchestrator.outbox_events,
            run=result.run,
            item=item,
            receipt=receipt,
            trace_id=_id("mismatch-trace"),
            causation_id=_id("mismatch-causation"),
        )

    assert "hello" not in str(captured.value)
    assert len(outbox.events) == event_count


@pytest.mark.parametrize(
    "forbidden_field",
    ("filename", "path", "parsed_text", "credentials", "provider_error"),
)
def test_projection_request_payload_rejects_content_bearing_fields(
    forbidden_field: str,
) -> None:
    values = {
        "ingestion_run_id": _id("safe-run"),
        "ingestion_run_source_id": _id("safe-item"),
        "source_revision_id": _id("safe-revision"),
        "processing_receipt_id": _id("safe-receipt"),
        "processor_kind": "document-ingestion",
        "parser_version": "text-markdown:v1",
        "parser_config_reference": "text-markdown-default:v1",
        "locator_contract_version": "physical-line-codepoint:v1",
        forbidden_field: "protected customer content",
    }

    with pytest.raises(ValidationError, match="Extra inputs are not permitted"):
        ProjectionRequestPayload.model_validate(values)
from spine.auth import AuthorizedRequestContext, RequestedAuthorizationScope
