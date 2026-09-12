"""Load local files into plain-text documents for Cognee ingest."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

SUPPORTED_SUFFIXES = {".md", ".txt", ".xlsx", ".pdf", ".docx"}


@dataclass
class Document:
    source_path: str
    title: str
    text: str
    doc_type: str  # markdown | text | bitrix_task | pdf | docx


def iter_files(root: Path) -> list[Path]:
    if root.is_file():
        return [root]
    files: list[Path] = []
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            if path.name.startswith("."):
                continue
            files.append(path)
    return files


def load_path(root: Path) -> list[Document]:
    docs: list[Document] = []
    for path in iter_files(root):
        suffix = path.suffix.lower()
        if suffix in {".md", ".txt"}:
            docs.append(_load_text(path))
        elif suffix == ".xlsx":
            docs.extend(_load_xlsx(path))
        elif suffix == ".pdf":
            docs.append(_load_pdf(path))
        elif suffix == ".docx":
            docs.append(_load_docx(path))
    return [d for d in docs if d.text.strip()]


def documents_to_remember_payloads(docs: list[Document]) -> list[str]:
    """Wrap each document with provenance so the graph keeps the source path."""
    payloads: list[str] = []
    for doc in docs:
        payloads.append(
            f"Source: {doc.source_path}\n"
            f"Type: {doc.doc_type}\n"
            f"Title: {doc.title}\n\n"
            f"{doc.text.strip()}"
        )
    return payloads


def _load_text(path: Path) -> Document:
    text = path.read_text(encoding="utf-8", errors="replace")
    return Document(
        source_path=str(path),
        title=path.stem,
        text=text,
        doc_type="markdown" if path.suffix.lower() == ".md" else "text",
    )


def _load_pdf(path: Path) -> Document:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        text = "\n".join(parts)
    except Exception as exc:  # noqa: BLE001
        text = f"(failed to extract PDF text: {exc})"
    return Document(
        source_path=str(path),
        title=path.stem,
        text=text,
        doc_type="pdf",
    )


def _load_docx(path: Path) -> Document:
    try:
        from docx import Document as DocxDocument

        document = DocxDocument(str(path))
        parts = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [_cell_str(c.text) for c in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    parts.append(line)
        text = "\n".join(parts)
    except Exception as exc:  # noqa: BLE001
        text = f"(failed to extract DOCX text: {exc})"
    return Document(
        source_path=str(path),
        title=path.stem,
        text=text,
        doc_type="docx",
    )


def _load_xlsx(path: Path) -> list[Document]:
    """Treat each non-empty row of a Bitrix-like export as a task card."""
    from openpyxl import load_workbook

    # Some Bitrix/YouTrack exports break under read_only=True (only first column
    # is visible). Prefer regular mode for correctness on small/medium sheets.
    wb = load_workbook(path, read_only=False, data_only=True)
    docs: list[Document] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = [
            tuple(cell.value for cell in row)
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)
        ]
        if not rows:
            continue
        header = [_cell_str(c) for c in rows[0]]
        data_rows = rows[1:] if any(header) else rows
        columns = header if any(header) else [f"col_{i}" for i in range(len(rows[0]))]
        for idx, row in enumerate(data_rows, start=1):
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            fields: list[str] = []
            for col_name, value in zip(columns, row):
                val = _cell_str(value)
                if not val:
                    continue
                label = col_name or "field"
                fields.append(f"{label}: {val}")
            if not fields:
                continue
            title = _guess_title(columns, row) or f"{path.stem}#{idx}"
            docs.append(
                Document(
                    source_path=f"{path}#{sheet_name}!row{idx}",
                    title=title,
                    text="\n".join(fields),
                    doc_type="bitrix_task",
                )
            )
    return docs


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _guess_title(columns: list[str], row: tuple) -> str:
    preferred = ("название", "title", "тема", "task", "задача", "name")
    lower_cols = [c.lower() for c in columns]
    for key in preferred:
        for i, col in enumerate(lower_cols):
            if key in col and i < len(row):
                val = _cell_str(row[i])
                if val:
                    return val[:120]
    # fallback: first non-empty cell
    for value in row:
        val = _cell_str(value)
        if val:
            return val[:120]
    return ""
