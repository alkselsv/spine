"""Strict, deterministic parsers for production document artifacts."""

from __future__ import annotations

from dataclasses import dataclass
import hashlib
from importlib.metadata import version
from io import BytesIO
import re
from typing import Final, Literal
from uuid import UUID
from zipfile import ZipFile

from spine.domain.artifacts import (
    ExtractedOoxmlText,
    ExtractedPageText,
    OoxmlBlockPathStep,
    OoxmlTextLocator,
    ParseCoverage,
    ParseExclusion,
    PageTextLocator,
    ParsedDocument,
    ParsedSegment,
    PhysicalLineLocator,
    XlsxCellLocator,
)
from spine.domain.sources import RevisionKind, SourceRevision, assert_revision_digest

LOCATOR_CONTRACT_VERSION: Final[str] = "physical-line-codepoint:v1"
PARSER_VERSION: Final[str] = "text-markdown:v1"
PARSER_CONFIG_REFERENCE: Final[str] = "text-markdown-default:v1"
PDF_LOCATOR_CONTRACT_VERSION: Final[str] = "pdf-page-extracted-codepoint:v1"
PDF_PARSER_VERSION: Final[str] = f"pdf:v1-pypdf-{version('pypdf')}"
PDF_EXTRACTION_MODE: Final[str] = "plain"
XLSX_LOCATOR_CONTRACT_VERSION: Final[str] = "xlsx-cell:v1"
XLSX_PARSER_VERSION: Final[str] = "xlsx-openpyxl:v1"
XLSX_PARSER_CONFIG_REFERENCE: Final[str] = "xlsx-openpyxl-default:v1"
XLSX_MEDIA_TYPE: Final[str] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MEDIA_TYPE: Final[str] = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DOCX_LOCATOR_CONTRACT_VERSION: Final[str] = "docx-ooxml-block-codepoint:v1"
DOCX_PARSER_VERSION: Final[str] = f"docx:v1-python-docx-{version('python-docx')}"
_WORDPROCESSINGML: Final[str] = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_UNSUPPORTED_DOCX_STORIES: Final[frozenset[str]] = frozenset(
    {"header", "footer", "footnotes", "endnotes", "comments"}
)
_UNSUPPORTED_DRAWING_TAGS: Final[frozenset[str]] = frozenset(
    {
        _WORDPROCESSINGML + "drawing",
        _WORDPROCESSINGML + "object",
        _WORDPROCESSINGML + "pict",
    }
)
_HEADING = re.compile(r"^(#{1,6})[ \t]+(.+?)\s*#*\s*$")


@dataclass(frozen=True, slots=True)
class TextParser:
    max_bytes: int = 10_000_000
    max_lines: int = 1_000_000

    def __post_init__(self) -> None:
        if self.max_bytes <= 0 or self.max_lines <= 0:
            raise ValueError("parser bounds must be positive")

    def parse(self, source_revision: SourceRevision, content: bytes, *, media_type: str) -> ParsedDocument:
        return _parse(source_revision, content, media_type=media_type, markdown=False, limits=self)


@dataclass(frozen=True, slots=True)
class MarkdownParser(TextParser):
    def parse(self, source_revision: SourceRevision, content: bytes, *, media_type: str = "text/markdown") -> ParsedDocument:
        return _parse(source_revision, content, media_type=media_type, markdown=True, limits=self)


@dataclass(frozen=True, slots=True)
class PdfParser:
    """Deterministic PDF text adapter with page-scoped extracted-text locators."""

    max_bytes: int = 10_000_000
    max_pages: int = 10_000
    max_page_chars: int = 2_000_000
    max_total_chars: int = 10_000_000

    def __post_init__(self) -> None:
        if min(self.max_bytes, self.max_pages, self.max_page_chars, self.max_total_chars) <= 0:
            raise ValueError("parser bounds must be positive")

    @property
    def config_reference(self) -> str:
        limits = (
            f"strict=true:mode={PDF_EXTRACTION_MODE}:"
            f"{self.max_bytes}:{self.max_pages}:"
            f"{self.max_page_chars}:{self.max_total_chars}"
        )
        digest = hashlib.sha256(limits.encode("ascii")).hexdigest()
        return f"pdf-pypdf-config:v1:{digest}"

    def parse(
        self,
        source_revision: SourceRevision,
        content: bytes,
        *,
        media_type: str = "application/pdf",
    ) -> ParsedDocument:
        _validate_binary_source(source_revision, content)

        if len(content) > self.max_bytes:
            return _pdf_result(
                source_revision,
                content,
                media_type=media_type,
                parser_config_reference=self.config_reference,
                segments=(),
                exclusions=(ParseExclusion(code="bounded_input", count=1),),
                extracted_length=0,
            )

        if not content.startswith(b"%PDF-"):
            return _pdf_result(
                source_revision,
                content,
                media_type=media_type,
                parser_config_reference=self.config_reference,
                segments=(),
                exclusions=(ParseExclusion(code="malformed_document", count=1),),
                extracted_length=0,
            )

        from pypdf import PdfReader

        try:
            reader = PdfReader(BytesIO(content), strict=True)
        except Exception:
            return _pdf_result(
                source_revision,
                content,
                media_type=media_type,
                parser_config_reference=self.config_reference,
                segments=(),
                exclusions=(ParseExclusion(code="malformed_document", count=1),),
                extracted_length=0,
            )
        segments: list[ParsedSegment] = []
        exclusions: list[ParseExclusion] = []
        extracted_length = 0
        try:
            page_count = len(reader.pages)
            if page_count == 0:
                exclusions.append(ParseExclusion(code="malformed_document", count=1))
            if page_count > self.max_pages:
                exclusions.append(
                    ParseExclusion(code="bounded_input", count=page_count - self.max_pages)
                )
            for page_number in range(1, min(page_count, self.max_pages) + 1):
                page = reader.pages[page_number - 1]
                try:
                    page_text = page.extract_text(extraction_mode=PDF_EXTRACTION_MODE) or ""
                except Exception:
                    exclusions.append(
                        ParseExclusion(
                            code="extraction_failed",
                            count=1,
                            locator=PageTextLocator(
                                page_number=page_number,
                                start_codepoint=0,
                                end_codepoint=0,
                            ),
                        )
                    )
                    continue
                if not isinstance(page_text, str):
                    exclusions.append(
                        ParseExclusion(
                            code="unlocatable_content",
                            count=1,
                            locator=PageTextLocator(
                                page_number=page_number,
                                start_codepoint=0,
                                end_codepoint=0,
                            ),
                        )
                    )
                    continue
                if not _has_exact_unicode_text(page_text):
                    exclusions.append(
                        ParseExclusion(
                            code="unlocatable_content",
                            count=1,
                            locator=PageTextLocator(
                                page_number=page_number,
                                start_codepoint=0,
                                end_codepoint=0,
                            ),
                        )
                    )
                    continue
                if len(page_text) > self.max_page_chars:
                    exclusions.append(
                        ParseExclusion(
                            code="bounded_input",
                            count=1,
                            locator=PageTextLocator(
                                page_number=page_number,
                                start_codepoint=0,
                                end_codepoint=0,
                            ),
                        )
                    )
                    continue
                if extracted_length + len(page_text) > self.max_total_chars:
                    exclusions.append(
                        ParseExclusion(
                            code="bounded_input",
                            count=min(page_count, self.max_pages) - page_number + 1,
                        )
                    )
                    break
                extracted_length += len(page_text)
                if not page_text.strip():
                    exclusions.append(
                        ParseExclusion(
                            code="blank_page",
                            count=1,
                            locator=PageTextLocator(
                                page_number=page_number,
                                start_codepoint=0,
                                end_codepoint=0,
                            ),
                        )
                    )
                    continue
                offset = 0
                for raw_line in page_text.splitlines(keepends=True):
                    text = raw_line.rstrip("\r\n")
                    locator = PageTextLocator(
                        page_number=page_number,
                        start_codepoint=offset,
                        end_codepoint=offset + len(text),
                    )
                    offset += len(raw_line)
                    segments.append(
                        ParsedSegment.model_construct(
                            segment_id=UUID(int=1),
                            segment_digest="0" * 64,
                            artifact_digest="0" * 64,
                            text=text,
                            text_digest=hashlib.sha256(text.encode("utf-8")).hexdigest(),
                            locator=locator,
                            heading_path=(),
                        )
                    )
        except Exception:
            exclusions.append(ParseExclusion(code="extraction_failed", count=1))
        return _pdf_result(
            source_revision,
            content,
            media_type=media_type,
            parser_config_reference=self.config_reference,
            segments=tuple(segments),
            exclusions=tuple(exclusions),
            extracted_length=extracted_length,
        )

    def resolve_locator(
        self,
        source_revision: SourceRevision,
        content: bytes,
        document: ParsedDocument,
        locator: PageTextLocator,
    ) -> str:
        """Reopen one pinned PDF revision and resolve an exact page-text range."""

        _validate_binary_source(source_revision, content)
        _assert_parser_identity(
            source_revision,
            document,
            parser_name="pdf",
            parser_version=PDF_PARSER_VERSION,
            parser_config_reference=self.config_reference,
            locator_contract_version=PDF_LOCATOR_CONTRACT_VERSION,
            media_type="application/pdf",
        )

        from pypdf import PdfReader

        try:
            reader = PdfReader(BytesIO(content), strict=True)
            page = reader.pages[locator.page_number - 1]
            page_text = page.extract_text(extraction_mode=PDF_EXTRACTION_MODE) or ""
            if not isinstance(page_text, str) or not _has_exact_unicode_text(page_text):
                raise ValueError("locator does not resolve")
            return document.resolve_locator(
                ExtractedPageText(page_number=locator.page_number, text=page_text),
                locator,
            )
        except (IndexError, TypeError, ValueError):
            raise ValueError("locator does not resolve") from None
        except Exception:
            raise ValueError("locator does not resolve") from None


@dataclass(slots=True)
class _DocxTraversalBudget:
    max_blocks: int
    max_paragraph_chars: int
    max_total_chars: int
    block_count: int = 0
    total_chars: int = 0
    blocks_exhausted: bool = False

    def admit_block(self) -> bool:
        if self.block_count >= self.max_blocks:
            self.blocks_exhausted = True
            return False
        self.block_count += 1
        return True

    def admit_text(self, text: str) -> bool:
        if len(text) > self.max_paragraph_chars:
            return False
        if self.total_chars + len(text) > self.max_total_chars:
            return False
        self.total_chars += len(text)
        return True


@dataclass(frozen=True, slots=True)
class DocxParser:
    """Deterministic WordprocessingML adapter with structural paragraph locators."""

    max_bytes: int = 10_000_000
    max_archive_entries: int = 10_000
    max_uncompressed_bytes: int = 100_000_000
    max_nesting_depth: int = 32
    max_blocks: int = 1_000_000
    max_paragraph_chars: int = 2_000_000
    max_total_chars: int = 10_000_000

    def __post_init__(self) -> None:
        if min(
            self.max_bytes,
            self.max_archive_entries,
            self.max_uncompressed_bytes,
            self.max_nesting_depth,
            self.max_blocks,
            self.max_paragraph_chars,
            self.max_total_chars,
        ) <= 0:
            raise ValueError("parser bounds must be positive")

    @property
    def config_reference(self) -> str:
        limits = (
            f"strict=true:{self.max_bytes}:"
            f"{self.max_archive_entries}:{self.max_uncompressed_bytes}:"
            f"{self.max_nesting_depth}:{self.max_blocks}:"
            f"{self.max_paragraph_chars}:{self.max_total_chars}"
        )
        digest = hashlib.sha256(limits.encode("ascii")).hexdigest()
        return f"docx-python-docx-config:v1:{digest}"

    def parse(
        self,
        source_revision: SourceRevision,
        content: bytes,
        *,
        media_type: str = DOCX_MEDIA_TYPE,
    ) -> ParsedDocument:
        _validate_binary_source(source_revision, content)

        if len(content) > self.max_bytes:
            return _docx_result(
                source_revision,
                content,
                media_type=media_type,
                parser_config_reference=self.config_reference,
                segments=(),
                exclusions=(ParseExclusion(code="bounded_input", count=1),),
            )

        try:
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
        except Exception:
            return _docx_result(
                source_revision,
                content,
                media_type=media_type,
                parser_config_reference=self.config_reference,
                segments=(),
                exclusions=(ParseExclusion(code="malformed_document", count=1),),
            )
        if (
            len(entries) > self.max_archive_entries
            or sum(entry.file_size for entry in entries) > self.max_uncompressed_bytes
        ):
            return _docx_result(
                source_revision,
                content,
                media_type=media_type,
                parser_config_reference=self.config_reference,
                segments=(),
                exclusions=(ParseExclusion(code="bounded_input", count=1),),
            )

        from docx import Document

        try:
            document = Document(BytesIO(content))
        except Exception:
            return _docx_result(
                source_revision,
                content,
                media_type=media_type,
                parser_config_reference=self.config_reference,
                segments=(),
                exclusions=(ParseExclusion(code="malformed_document", count=1),),
            )
        segments, structural_exclusions = _walk_ooxml_blocks(
            document.element.body,
            budget=_DocxTraversalBudget(
                max_blocks=self.max_blocks,
                max_paragraph_chars=self.max_paragraph_chars,
                max_total_chars=self.max_total_chars,
            ),
            max_nesting_depth=self.max_nesting_depth,
        )
        unsupported_story_count = sum(
            relationship.reltype.rsplit("/", 1)[-1] in _UNSUPPORTED_DOCX_STORIES
            for relationship in document.part.rels.values()
        )
        exclusions = structural_exclusions
        if unsupported_story_count:
            exclusions += (ParseExclusion(code="unsupported_story", count=unsupported_story_count),)
        return _docx_result(
            source_revision,
            content,
            media_type=media_type,
            parser_config_reference=self.config_reference,
            segments=segments,
            exclusions=exclusions,
        )

    def resolve_locator(
        self,
        source_revision: SourceRevision,
        content: bytes,
        document: ParsedDocument,
        locator: OoxmlTextLocator,
    ) -> str:
        """Reopen one pinned DOCX revision and resolve an exact OOXML range."""

        _validate_binary_source(source_revision, content)
        _assert_parser_identity(
            source_revision,
            document,
            parser_name="docx",
            parser_version=DOCX_PARSER_VERSION,
            parser_config_reference=self.config_reference,
            locator_contract_version=DOCX_LOCATOR_CONTRACT_VERSION,
            media_type=DOCX_MEDIA_TYPE,
        )

        from docx import Document

        try:
            source_document = Document(BytesIO(content))
            paragraph = _resolve_ooxml_paragraph(
                source_document.element.body,
                locator.block_path,
            )
            paragraph_text = _ooxml_paragraph_text(paragraph)
            if locator.end_codepoint > len(paragraph_text):
                raise ValueError("locator does not resolve")
            return document.resolve_locator(
                ExtractedOoxmlText(
                    block_path=locator.block_path,
                    text=paragraph_text,
                ),
                locator,
            )
        except (IndexError, StopIteration, TypeError, ValueError):
            raise ValueError("locator does not resolve") from None
        except Exception:
            raise ValueError("locator does not resolve") from None


@dataclass(frozen=True, slots=True)
class XlsxParser:
    """Parse only safe, citation-grade values from an authorized XLSX original."""

    max_bytes: int = 50_000_000
    max_cells: int = 1_000_000
    max_worksheets: int = 256

    def __post_init__(self) -> None:
        if self.max_bytes <= 0 or self.max_cells <= 0 or self.max_worksheets <= 0:
            raise ValueError("parser bounds must be positive")

    def parse(
        self,
        source_revision: SourceRevision,
        content: bytes,
        *,
        media_type: str = XLSX_MEDIA_TYPE,
    ) -> ParsedDocument:
        return _parse_xlsx(source_revision, content, media_type=media_type, limits=self)

    def resolve_locator(
        self, document: ParsedDocument, content: bytes, locator: XlsxCellLocator
    ) -> str:
        return resolve_xlsx_locator(document, content, locator)


def parse_document(
    source_revision: SourceRevision,
    content: bytes,
    *,
    media_type: str | None = None,
    max_bytes: int = 10_000_000,
    max_lines: int = 1_000_000,
    max_cells: int = 1_000_000,
    max_worksheets: int = 256,
) -> ParsedDocument:
    """Parse an immutable original with the selected production adapter."""

    if source_revision.kind is not RevisionKind.CONTENT:
        raise ValueError("only content revisions can be parsed")
    assert_revision_digest(source_revision)
    selected_media_type = (media_type or source_revision.media_type or "").lower()
    if selected_media_type == XLSX_MEDIA_TYPE:
        return XlsxParser(
            max_bytes=max_bytes, max_cells=max_cells, max_worksheets=max_worksheets
        ).parse(source_revision, content, media_type=XLSX_MEDIA_TYPE)
    if selected_media_type in {"text/markdown", "text/x-markdown"}:
        return MarkdownParser(max_bytes=max_bytes, max_lines=max_lines).parse(
            source_revision, content, media_type="text/markdown"
        )
    if selected_media_type == "text/plain":
        return TextParser(max_bytes=max_bytes, max_lines=max_lines).parse(
            source_revision, content, media_type="text/plain"
        )
    if selected_media_type == "application/pdf":
        return PdfParser(max_bytes=max_bytes).parse(source_revision, content)
    if selected_media_type == DOCX_MEDIA_TYPE:
        return DocxParser(max_bytes=max_bytes).parse(source_revision, content)
    raise ValueError("unsupported parser media type")


def _parse(
    source_revision: SourceRevision,
    content: bytes,
    *,
    media_type: str,
    markdown: bool,
    limits: TextParser,
) -> ParsedDocument:
    if len(content) != source_revision.byte_length or hashlib.sha256(content).hexdigest() != source_revision.original_sha256:
        raise ValueError("original content does not match source revision")
    if len(content) > limits.max_bytes:
        return _failed_document(source_revision, content, media_type=media_type, markdown=markdown,
                                 code="bounded_input")
    try:
        text = content.decode("utf-8", errors="strict")
    except UnicodeDecodeError:
        return _failed_document(source_revision, content, media_type=media_type, markdown=markdown,
                                 code="invalid_utf8")

    lines = list(_lines(text))
    exclusions: list[ParseExclusion] = []
    if len(lines) > limits.max_lines:
        exclusions.append(ParseExclusion(code="bounded_input", count=len(lines) - limits.max_lines))
        lines = lines[:limits.max_lines]
    segments: list[ParsedSegment] = []
    headings: list[str] = []
    fence_marker: str | None = None
    unsupported_lines = 0
    for line_number, start, raw_line in lines:
        line = raw_line.rstrip("\r\n")
        heading_path = ()
        if markdown:
            stripped = line.lstrip()
            if fence_marker is not None:
                unsupported_lines += 1
                if stripped.startswith(fence_marker):
                    fence_marker = None
                continue
            if stripped.startswith("```") or stripped.startswith("~~~"):
                fence_marker = stripped[:3]
                unsupported_lines += 1
                continue
            match = _HEADING.match(line)
            if match:
                level = len(match.group(1))
                title = match.group(2).strip()
                if not title:
                    exclusions.append(ParseExclusion(code="malformed_structure", count=1,
                        locator=PhysicalLineLocator(line_number=line_number, start_codepoint=start, end_codepoint=start + len(line))))
                else:
                    headings = headings[: level - 1] + [title]
                heading_path = tuple(headings)
            else:
                heading_path = tuple(headings)
        locator = PhysicalLineLocator(
            line_number=line_number,
            start_codepoint=start,
            end_codepoint=start + len(line),
        )
        text_digest = hashlib.sha256(line.encode("utf-8")).hexdigest()
        # The document digest is the parent identity, so segment identities are
        # completed by ParsedDocument.from_source after the parent is derived.
        segments.append(ParsedSegment.model_construct(
            segment_id=UUID(int=1), segment_digest="0" * 64,
            artifact_digest="0" * 64, text=line, text_digest=text_digest,
            locator=locator, heading_path=heading_path,
        ))
    if fence_marker is not None:
        exclusions.append(ParseExclusion(code="malformed_structure", count=1))
    if unsupported_lines:
        exclusions.append(ParseExclusion(code="unsupported_structure", count=unsupported_lines))
    coverage = ParseCoverage.PARTIAL if exclusions else ParseCoverage.COMPLETE
    return ParsedDocument.from_source(
        source_revision, parser_name="markdown" if markdown else "text", parser_version=PARSER_VERSION,
        parser_config_reference=PARSER_CONFIG_REFERENCE, locator_contract_version=LOCATOR_CONTRACT_VERSION,
        media_type=media_type, coverage=coverage, segments=tuple(segments), exclusions=tuple(exclusions),
        source_byte_length=len(content), source_codepoint_length=len(text),
    )


def _lines(text: str) -> list[tuple[int, int, str]]:
    result: list[tuple[int, int, str]] = []
    offset = 0
    for line_number, line in enumerate(text.splitlines(keepends=True), start=1):
        result.append((line_number, offset, line))
        offset += len(line)
    if text and (not result or offset < len(text)):
        result.append((len(result) + 1, offset, text[offset:]))
    return result


def _failed_document(
    source_revision: SourceRevision,
    content: bytes,
    *,
    media_type: str,
    markdown: bool,
    code: Literal["bounded_input", "invalid_utf8"],
) -> ParsedDocument:
    return ParsedDocument.from_source(
        source_revision,
        parser_name="markdown" if markdown else "text",
        parser_version=PARSER_VERSION,
        parser_config_reference=PARSER_CONFIG_REFERENCE,
        locator_contract_version=LOCATOR_CONTRACT_VERSION,
        media_type=media_type,
        coverage=ParseCoverage.FAILED,
        segments=(),
        exclusions=(ParseExclusion(code=code, count=1),),
        source_byte_length=len(content),
        source_codepoint_length=0,
    )


def _pdf_result(
    source_revision: SourceRevision,
    content: bytes,
    *,
    media_type: str,
    parser_config_reference: str,
    segments: tuple[ParsedSegment, ...],
    exclusions: tuple[ParseExclusion, ...],
    extracted_length: int,
) -> ParsedDocument:
    coverage = _coverage_for(segments, exclusions)
    return ParsedDocument.from_source(
        source_revision,
        parser_name="pdf",
        parser_version=PDF_PARSER_VERSION,
        parser_config_reference=parser_config_reference,
        locator_contract_version=PDF_LOCATOR_CONTRACT_VERSION,
        media_type=media_type,
        coverage=coverage,
        segments=segments,
        exclusions=exclusions,
        source_byte_length=len(content),
        source_codepoint_length=extracted_length,
    )


def _assert_parser_identity(
    source_revision: SourceRevision,
    document: ParsedDocument,
    *,
    parser_name: str,
    parser_version: str,
    parser_config_reference: str,
    locator_contract_version: str,
    media_type: str,
) -> None:
    expected = (
        source_revision.workspace_id,
        source_revision.environment_id,
        source_revision.source_object_id,
        source_revision.revision_id,
        source_revision.revision_digest,
        source_revision.original_sha256,
        parser_name,
        parser_version,
        parser_config_reference,
        locator_contract_version,
        media_type,
    )
    actual = (
        document.workspace_id,
        document.environment_id,
        document.source_object_id,
        document.source_revision_id,
        document.source_revision_digest,
        document.original_sha256,
        document.parser_name,
        document.parser_version,
        document.parser_config_reference,
        document.locator_contract_version,
        document.media_type,
    )
    if actual != expected:
        raise ValueError("locator does not resolve")


def _validate_binary_source(source_revision: SourceRevision, content: bytes) -> None:
    if source_revision.kind is not RevisionKind.CONTENT:
        raise ValueError("only content revisions can be parsed")
    assert_revision_digest(source_revision)
    if (
        len(content) != source_revision.byte_length
        or hashlib.sha256(content).hexdigest() != source_revision.original_sha256
    ):
        raise ValueError("original content does not match source revision")
def _walk_ooxml_blocks(
    parent: object,
    prefix: tuple[OoxmlBlockPathStep, ...] = (),
    segments: list[ParsedSegment] | None = None,
    exclusions: list[ParseExclusion] | None = None,
    *,
    budget: _DocxTraversalBudget | None = None,
    nesting_depth: int = 0,
    max_nesting_depth: int = 32,
) -> tuple[tuple[ParsedSegment, ...], tuple[ParseExclusion, ...]]:
    segment_values = [] if segments is None else segments
    exclusion_values = [] if exclusions is None else exclusions
    active_budget = budget or _DocxTraversalBudget(
        max_blocks=1_000_000,
        max_paragraph_chars=2_000_000,
        max_total_chars=10_000_000,
    )
    for block_index, child in enumerate(parent):  # type: ignore[arg-type]
        if active_budget.blocks_exhausted:
            break
        if child.tag in {
            _WORDPROCESSINGML + "sectPr",
            _WORDPROCESSINGML + "tcPr",
        }:
            continue
        if not active_budget.admit_block():
            exclusion_values.append(ParseExclusion(code="bounded_input", count=1))
            break
        if child.tag == _WORDPROCESSINGML + "p":
            text = _ooxml_paragraph_text(child)
            locator = OoxmlTextLocator(
                block_path=prefix
                + (OoxmlBlockPathStep(kind="paragraph", index=block_index),),
                start_codepoint=0,
                end_codepoint=len(text),
            )
            if text and active_budget.admit_text(text):
                segment_values.append(
                    ParsedSegment.model_construct(
                        segment_id=UUID(int=1),
                        segment_digest="0" * 64,
                        artifact_digest="0" * 64,
                        text=text,
                        text_digest=hashlib.sha256(text.encode("utf-8")).hexdigest(),
                        locator=locator,
                        heading_path=(),
                    )
                )
            elif text:
                exclusion_values.append(ParseExclusion(code="bounded_input", count=1))
            drawing_count = sum(
                node.tag in _UNSUPPORTED_DRAWING_TAGS for node in child.iter()
            )
            if drawing_count:
                exclusion_values.append(
                    ParseExclusion(
                        code="unsupported_drawing",
                        count=drawing_count,
                        locator=locator,
                    )
                )
        elif child.tag == _WORDPROCESSINGML + "tbl":
            if nesting_depth >= max_nesting_depth:
                exclusion_values.append(ParseExclusion(code="bounded_input", count=1))
                continue
            table_path = prefix + (OoxmlBlockPathStep(kind="table", index=block_index),)
            row_index = 0
            for table_child in child:
                if active_budget.blocks_exhausted:
                    break
                if table_child.tag in {
                    _WORDPROCESSINGML + "tblPr",
                    _WORDPROCESSINGML + "tblGrid",
                }:
                    continue
                if not active_budget.admit_block():
                    exclusion_values.append(ParseExclusion(code="bounded_input", count=1))
                    break
                if table_child.tag != _WORDPROCESSINGML + "tr":
                    exclusion_values.append(
                        ParseExclusion(code="unsupported_structure", count=1)
                    )
                    continue
                row = table_child
                row_path = table_path + (OoxmlBlockPathStep(kind="row", index=row_index),)
                cell_index = 0
                for row_child in row:
                    if active_budget.blocks_exhausted:
                        break
                    if row_child.tag == _WORDPROCESSINGML + "trPr":
                        continue
                    if not active_budget.admit_block():
                        exclusion_values.append(ParseExclusion(code="bounded_input", count=1))
                        break
                    if row_child.tag != _WORDPROCESSINGML + "tc":
                        exclusion_values.append(
                            ParseExclusion(code="unsupported_structure", count=1)
                        )
                        continue
                    cell = row_child
                    if (
                        cell.find(
                            _WORDPROCESSINGML + "tcPr/" + _WORDPROCESSINGML + "gridSpan"
                        )
                        is not None
                        or cell.find(
                            _WORDPROCESSINGML + "tcPr/" + _WORDPROCESSINGML + "vMerge"
                        )
                        is not None
                    ):
                        exclusion_values.append(
                            ParseExclusion(code="ambiguous_merged_cell", count=1)
                        )
                        cell_index += 1
                        continue
                    cell_path = row_path + (OoxmlBlockPathStep(kind="cell", index=cell_index),)
                    _walk_ooxml_blocks(
                        cell,
                        cell_path,
                        segment_values,
                        exclusion_values,
                        budget=active_budget,
                        nesting_depth=nesting_depth + 1,
                        max_nesting_depth=max_nesting_depth,
                    )
                    cell_index += 1
                row_index += 1
        else:
            exclusion_values.append(ParseExclusion(code="unsupported_structure", count=1))
    return tuple(segment_values), tuple(exclusion_values)


def _resolve_ooxml_paragraph(
    body: object,
    block_path: tuple[OoxmlBlockPathStep, ...],
) -> object:
    parent = body
    position = 0
    while position < len(block_path):
        block = block_path[position]
        child = parent[block.index]  # type: ignore[index]
        if block.kind == "paragraph":
            if child.tag != _WORDPROCESSINGML + "p" or position != len(block_path) - 1:
                raise ValueError("locator does not resolve")
            return child
        if block.kind != "table" or child.tag != _WORDPROCESSINGML + "tbl":
            raise ValueError("locator does not resolve")
        row_step = block_path[position + 1]
        cell_step = block_path[position + 2]
        if row_step.kind != "row" or cell_step.kind != "cell":
            raise ValueError("locator does not resolve")
        row = child.findall(_WORDPROCESSINGML + "tr")[row_step.index]
        parent = row.findall(_WORDPROCESSINGML + "tc")[cell_step.index]
        position += 3
    raise ValueError("locator does not resolve")


def _ooxml_paragraph_text(paragraph: object) -> str:
    pieces: list[str] = []
    for node in paragraph.iter():  # type: ignore[attr-defined]
        if any(
            ancestor.tag in _UNSUPPORTED_DRAWING_TAGS
            for ancestor in node.iterancestors()
        ):
            continue
        if node.tag == _WORDPROCESSINGML + "t":
            pieces.append(node.text or "")
        elif node.tag == _WORDPROCESSINGML + "tab":
            pieces.append("\t")
        elif node.tag in {_WORDPROCESSINGML + "br", _WORDPROCESSINGML + "cr"}:
            pieces.append("\n")
    return "".join(pieces)


def _docx_result(
    source_revision: SourceRevision,
    content: bytes,
    *,
    media_type: str,
    parser_config_reference: str,
    segments: tuple[ParsedSegment, ...],
    exclusions: tuple[ParseExclusion, ...],
) -> ParsedDocument:
    coverage = _coverage_for(segments, exclusions)
    return ParsedDocument.from_source(
        source_revision,
        parser_name="docx",
        parser_version=DOCX_PARSER_VERSION,
        parser_config_reference=parser_config_reference,
        locator_contract_version=DOCX_LOCATOR_CONTRACT_VERSION,
        media_type=media_type,
        coverage=coverage,
        segments=segments,
        exclusions=exclusions,
        source_byte_length=len(content),
        source_codepoint_length=sum(len(segment.text) for segment in segments),
    )


def _coverage_for(
    segments: tuple[ParsedSegment, ...],
    exclusions: tuple[ParseExclusion, ...],
) -> ParseCoverage:
    if exclusions and not segments:
        return ParseCoverage.FAILED
    if exclusions:
        return ParseCoverage.PARTIAL
    return ParseCoverage.COMPLETE


def _has_exact_unicode_text(value: str) -> bool:
    if "\ufffd" in value or "\x00" in value:
        return False
    try:
        value.encode("utf-8", errors="strict")
    except UnicodeEncodeError:
        return False
    return True


def _parse_xlsx(
    source_revision: SourceRevision,
    content: bytes,
    *,
    media_type: str,
    limits: XlsxParser,
) -> ParsedDocument:
    if source_revision.kind is not RevisionKind.CONTENT:
        raise ValueError("only content revisions can be parsed")
    assert_revision_digest(source_revision)
    if (
        len(content) != source_revision.byte_length
        or hashlib.sha256(content).hexdigest() != source_revision.original_sha256
    ):
        raise ValueError("original content does not match source revision")
    if len(content) > limits.max_bytes:
        return _failed_xlsx_document(source_revision, content, media_type=media_type, code="bounded_input")

    try:
        from openpyxl import load_workbook

        values_workbook = load_workbook(
            BytesIO(content), read_only=False, data_only=True, keep_links=True
        )
        formulas_workbook = load_workbook(
            BytesIO(content), read_only=False, data_only=False, keep_links=True
        )
    except Exception:
        # The parser boundary is deliberately fail-closed for malformed ZIP/XML
        # and unsupported workbook structures without returning provider errors.
        return _failed_xlsx_document(source_revision, content, media_type=media_type, code="malformed_structure")

    try:
        exclusions: list[ParseExclusion] = []
        segments: list[ParsedSegment] = []
        if len(values_workbook.worksheets) > limits.max_worksheets:
            exclusions.append(
                ParseExclusion(
                    code="bounded_input", count=len(values_workbook.worksheets) - limits.max_worksheets
                )
            )
        worksheets = values_workbook.worksheets[: limits.max_worksheets]
        formula_sheets = {sheet.title: sheet for sheet in formulas_workbook.worksheets}
        external_links = len(getattr(formulas_workbook, "_external_links", ()))
        formula_without_cache = 0
        hidden_content = 0
        unsupported_structure = 0
        external_formula_count = external_links
        scanned_cells = 0
        bounded = False

        for sheet in worksheets:
            formula_sheet = formula_sheets.get(sheet.title)
            if formula_sheet is None:
                unsupported_structure += 1
                continue
            if sheet.sheet_state != "visible":
                hidden_content += 1
                continue
            unsupported_structure += len(getattr(sheet, "_charts", ()))
            unsupported_structure += len(getattr(sheet, "_images", ()))
            for row in sheet.iter_rows():
                for cell in row:
                    if scanned_cells >= limits.max_cells:
                        bounded = True
                        break
                    scanned_cells += 1
                    formula_cell = formula_sheet[cell.coordinate]
                    if _xlsx_cell_hidden(sheet, cell.row, cell.column):
                        if cell.value is not None:
                            hidden_content += 1
                        continue
                    value = cell.value
                    formula = formula_cell.value
                    locator = XlsxCellLocator(
                        sheet=sheet.title,
                        row=cell.row,
                        cell_range=_merged_cell_range(sheet, cell.coordinate),
                    )
                    if formula_cell.data_type == "f" and _is_external_formula(formula):
                        external_formula_count += 1
                        continue
                    if formula_cell.data_type == "f" and value is None:
                        formula_without_cache += 1
                        continue
                    if formula_cell.data_type == "e" or isinstance(
                        value, (bytes, bytearray, dict, list, tuple, set)
                    ):
                        unsupported_structure += 1
                        continue
                    if value is None:
                        continue
                    text = _xlsx_cell_text(value)
                    if not text.strip():
                        continue
                    segments.append(_placeholder_segment(text=text, locator=locator))
                if bounded:
                    break
            if bounded:
                break

        if bounded:
            exclusions.append(ParseExclusion(code="bounded_input", count=1))
        if formula_without_cache:
            exclusions.append(
                ParseExclusion(code="formula_without_cached_value", count=formula_without_cache)
            )
        if hidden_content:
            exclusions.append(ParseExclusion(code="hidden_content", count=hidden_content))
        if external_formula_count:
            exclusions.append(ParseExclusion(code="external_link", count=external_formula_count))
        if unsupported_structure:
            exclusions.append(ParseExclusion(code="unsupported_structure", count=unsupported_structure))
        if not segments and not exclusions:
            exclusions.append(ParseExclusion(code="unsupported_structure", count=1))
        coverage = ParseCoverage.COMPLETE if not exclusions else ParseCoverage.PARTIAL
        return ParsedDocument.from_source(
            source_revision,
            parser_name="xlsx",
            parser_version=XLSX_PARSER_VERSION,
            parser_config_reference=XLSX_PARSER_CONFIG_REFERENCE,
            locator_contract_version=XLSX_LOCATOR_CONTRACT_VERSION,
            media_type=media_type,
            coverage=coverage,
            segments=tuple(segments),
            exclusions=tuple(exclusions),
            source_byte_length=len(content),
            source_codepoint_length=0,
        )
    finally:
        values_workbook.close()
        formulas_workbook.close()


def resolve_xlsx_locator(
    document: ParsedDocument, content: bytes, locator: XlsxCellLocator
) -> str:
    """Reopen one exact XLSX revision and resolve a persisted cell locator."""

    if document.parser_name != "xlsx" or document.parser_version != XLSX_PARSER_VERSION:
        raise ValueError("locator parser identity does not match")
    if document.locator_contract_version != XLSX_LOCATOR_CONTRACT_VERSION:
        raise ValueError("locator contract identity does not match")
    if (
        len(content) != document.source_byte_length
        or hashlib.sha256(content).hexdigest() != document.original_sha256
    ):
        raise ValueError("original content does not match source revision")
    segment = next((item for item in document.segments if item.locator == locator), None)
    if segment is None:
        raise ValueError("locator does not resolve")
    workbook = None
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=False, data_only=True, keep_links=True)
        sheet = workbook[locator.sheet]
        start_coordinate = locator.cell_range.split(":", 1)[0]
        cell = sheet[start_coordinate]
        if cell.row != locator.row or cell.coordinate != start_coordinate:
            raise ValueError("locator does not resolve")
        if locator.cell_range not in {str(value) for value in sheet.merged_cells.ranges} and ":" in locator.cell_range:
            raise ValueError("locator does not resolve")
        value = cell.value
    except Exception:
        raise ValueError("locator does not resolve") from None
    finally:
        if workbook is not None:
            workbook.close()
    if value is None or _xlsx_cell_text(value) != segment.text:
        raise ValueError("locator does not resolve")
    return segment.text


def _placeholder_segment(*, text: str, locator: XlsxCellLocator) -> ParsedSegment:
    text_digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    return ParsedSegment.model_construct(
        segment_id=UUID(int=1),
        segment_digest="0" * 64,
        artifact_digest="0" * 64,
        text=text,
        text_digest=text_digest,
        locator=locator,
        heading_path=(),
    )


def _failed_xlsx_document(
    source_revision: SourceRevision,
    content: bytes,
    *,
    media_type: str,
    code: Literal["bounded_input", "malformed_structure"],
) -> ParsedDocument:
    return ParsedDocument.from_source(
        source_revision,
        parser_name="xlsx",
        parser_version=XLSX_PARSER_VERSION,
        parser_config_reference=XLSX_PARSER_CONFIG_REFERENCE,
        locator_contract_version=XLSX_LOCATOR_CONTRACT_VERSION,
        media_type=media_type,
        coverage=ParseCoverage.FAILED,
        segments=(),
        exclusions=(ParseExclusion(code=code, count=1),),
        source_byte_length=len(content),
        source_codepoint_length=0,
    )


def _merged_cell_range(sheet: object, coordinate: str) -> str:
    for merged_range in getattr(sheet, "merged_cells").ranges:
        if coordinate in merged_range:
            return str(merged_range)
    return coordinate


def _xlsx_cell_hidden(sheet: object, row: int, column: int) -> bool:
    row_dimension = getattr(sheet, "row_dimensions").get(row)
    if row_dimension is not None and row_dimension.hidden:
        return True
    column_dimensions = getattr(sheet, "column_dimensions")
    for dimension in column_dimensions.values():
        if not dimension.hidden:
            continue
        minimum = dimension.min or column
        maximum = dimension.max or column
        if minimum <= column <= maximum:
            return True
    return False


def _is_external_formula(value: object) -> bool:
    return isinstance(value, str) and "[" in value and "]" in value


def _xlsx_cell_text(value: object) -> str:
    return str(value)
